"""Build manuscript insertion material and exact data-derived600dpi figures."""
from pathlib import Path
import argparse
import json
import html
import shutil
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,Mm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer,PageBreak,Image as PDFImage
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
import make_extension_common as c


LABELS={
    'PAD_baseline':'PAD-only transfer', 'target_weighted':'V-VI addition, weighted',
    'target_unweighted':'V-VI addition, unweighted', 'lighter_weighted':'Outcome-matched lighter, weighted',
    'all_DDI_unweighted':'All-DDI addition, unweighted', 'all_DDI_source_weighted':'All-DDI addition, weighted',
    'all_DDI_group_reweighted':'All-DDI addition, tone reweighted', 'all_DDI_budget_matched':'All-DDI, matched annotation budget',
    'DDI_only_class_balanced':'DDI-only, class balanced', 'DDI_only_group_balanced':'DDI-only, six groups balanced',
    'DFR_style':'DFR-style last-layer refit', 'target_budget_25':'25 V-VI images',
    'target_budget_50':'50 V-VI images','target_budget_100':'100 V-VI images',
    'diagnosis_matched_target':'Diagnosis-matched V-VI','diagnosis_matched_lighter':'Diagnosis-matched lighter',
}


def fmt(point,lo,hi):
    return f'{point:.3f} [{lo:.3f}, {hi:.3f}]'


def build(output):
    out=Path(output);m=pd.read_csv(out/'results/all_arm_metrics.csv');co=pd.read_csv(out/'results/all_paired_contrasts.csv')
    v=m[m.cohort=='V-VI'].set_index('arm');cv=co[(co.cohort=='V-VI')&(co.metric=='AUROC')]
    e1=['PAD_baseline','target_weighted','target_unweighted','lighter_weighted','all_DDI_unweighted',
        'all_DDI_source_weighted','all_DDI_group_reweighted','all_DDI_budget_matched','DDI_only_class_balanced',
        'DDI_only_group_balanced','DFR_style']
    figdir=out/'figures';figdir.mkdir(exist_ok=True)
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'axes.spines.top':False,'axes.spines.right':False,
                         'savefig.facecolor':'white','axes.labelcolor':'#172334','text.color':'#172334'})
    figures=[]
    def save(fig,name,source,filters,description):
        fig.savefig(figdir/(name+'.png'),dpi=600,bbox_inches='tight',metadata={'Software':'matplotlib; original numeric output only'})
        fig.savefig(figdir/(name+'.pdf'),bbox_inches='tight');plt.close(fig)
        with Image.open(figdir/(name+'.png')) as im:size=im.size;dpi=im.info.get('dpi');c.require(min(dpi)>599.9,'600dpi figure not preserved')
        figures.append({'figure':name,'source_csv':source,'source_sha256':c.sha(out/source),'row_filters':filters,
                        'plot_description':description,'generation_code_sha256':c.sha(__file__),
                        'pixels':list(size),'dpi':list(dpi),'PNG_sha256':c.sha(figdir/(name+'.png')),
                        'PDF_sha256':c.sha(figdir/(name+'.pdf')),'synthetic_or_generated_image_used':False})
    fig,ax=plt.subplots(figsize=(7.5,4.8))
    for i,arm in enumerate(e1):
        r=v.loc[arm];color='#1F5A99' if arm.startswith('target') else ('#208171' if arm.startswith('DDI_only') else '#75859A')
        ax.errorbar(r.AUROC,i,xerr=[[r.AUROC-r.AUROC_CI_low],[r.AUROC_CI_high-r.AUROC]],fmt='o',capsize=3,color=color,ms=5)
    ax.set_yticks(range(len(e1)),[LABELS[a] for a in e1]);ax.invert_yaxis();ax.axvline(.5,ls='--',lw=.8,color='#BBBBBB')
    ax.set_xlim(.36,.94);ax.set_xlabel('V-VI AUROC with exploratory paired 95% CI');ax.grid(axis='x',alpha=.17)
    ax.set_title('E1: domain adaptation and bias baselines',loc='left',fontweight='bold',pad=12)
    fig.tight_layout();save(fig,'Figure_E1_baseline_comparison','results/Table_E1_baselines.csv',{'cohort':'V-VI'},'Held-out image-fold AUROC; annotation budgets reported in TableE1, not equal for full-DDI controls')
    fig,ax=plt.subplots(figsize=(6.5,3.3));arms=['PAD_baseline','target_budget_25','target_budget_50','target_budget_100','target_weighted'];x=[0,25,50,100,165.4]
    y=[v.loc[a].AUROC for a in arms];low=[v.loc[a].AUROC_CI_low for a in arms];high=[v.loc[a].AUROC_CI_high for a in arms]
    ax.plot(x,y,'o-',color='#1F5A99',lw=1.7);ax.fill_between(x,low,high,color='#1F5A99',alpha=.12)
    ax.errorbar(x,y,yerr=[np.array(y)-low,np.array(high)-y],fmt='none',ecolor='#1F5A99',capsize=3)
    ax.set_xticks(x,['0','25','50','100','All\n165-166']);ax.set_ylim(.35,.92);ax.set_xlabel('V-VI images added in each non-test fold')
    ax.set_ylabel('V-VI AUROC');ax.set_title('E2: annotation learning curve',loc='left',fontweight='bold');ax.grid(alpha=.15)
    fig.tight_layout();save(fig,'Figure_E2_learning_curve','results/Table_E2_learning_curve.csv',{'cohort':'V-VI'},'20 nested sampled panels per size25/50/100; full165-166 uses one saved target arm; band is conditional image/panel CI')
    fig,ax=plt.subplots(figsize=(7.2,2.8));pairs=[('target_weighted','lighter_weighted'),('diagnosis_matched_target','diagnosis_matched_lighter')]
    plot_labels=['Outcome-matched controls\n165-166 images/source/fold','Exact diagnosis-overlap controls\n114-125 images/source/fold']
    for i,(left,right) in enumerate(pairs):
        r=cv[(cv.left_arm==left)&(cv.right_arm==right)].iloc[0]
        ax.errorbar(r.difference_unrounded,i,xerr=[[r.difference_unrounded-r.CI_low],[r.CI_high-r.difference_unrounded]],fmt='o',capsize=4,color='#1F5A99',ms=6)
    ax.axvline(0,color='#606B78',lw=.9,ls='--');ax.set_yticks([0,1],plot_labels);ax.invert_yaxis()
    ax.set_xlim(-.04,.34);ax.set_xlabel('V-VI AUROC difference: target minus lighter [95% CI]');ax.grid(axis='x',alpha=.15)
    ax.set_title('E3: disease-mix sensitivity',loc='left',fontweight='bold');fig.tight_layout()
    save(fig,'Figure_E3_diagnosis_sensitivity','results/all_paired_contrasts.csv',{'cohort':'V-VI','metric':'AUROC','pairs':pairs},'Different adaptation support and budgets; not identical causal-effect estimates; all same held-out VVI images')
    c.write_json(figdir/'figure_provenance.json',figures)
    coverage=pd.read_csv(out/'results/Table_E3_matching_coverage.csv');budget=pd.read_csv(out/'results/annotation_budgets.csv')
    def metric(a):
        r=v.loc[a];return fmt(r.AUROC,r.AUROC_CI_low,r.AUROC_CI_high)
    def delta(a,b):
        r=cv[(cv.left_arm==a)&(cv.right_arm==b)].iloc[0]
        return f'{r.difference_unrounded:+.3f} [{r.CI_low:+.3f}, {r.CI_high:+.3f}]'
    # Shared semantic content drives both DOCX and independent PDF rendering.
    pages=[]
    p1=[('title','Additional controlled experiments for the MAKE manuscript'),
        ('small','6 October2026 | Scientific insertion material | E1-E3 completed; E4 inputs pending'),
        ('p','This addendum reports a distinct exploratory extension of the verified PAD/DDI analysis. Its protocol was recorded before the additional fitting, after the original DDI results were known. It does not reconstruct the unavailable plan a2e188c2…ca4ac4 or replace the designated20261006 MDPI manuscripts and their32-reference bibliography.'),
        ('h','Methods suitable for insertion'),
        ('p','All experiments used the existing frozen MedSigLIP448 representation:2,268 curated PAD-UFES-20 images and656 DDI images, with1,152 L2-normalized features per image. The L2 logistic head retained C=10, selected previously using PAD only. DDI predictions used the same five outer image folds; every adaptation and internal tuning image was outside its corresponding test fold. Released DDI metadata do not establish patient-disjoint folds.'),
        ('p','E1 compared targeted V-VI addition against full-DDI addition, source and tone reweighting, and DDI-only class- or six-group-balanced heads. A20-panel all-tone control matched the targeted addition in total annotation budget and positive/negative counts. Source-weighted additions multiplied pooled inverse-class weights by the square root of the PAD-to-added-DDI sample-size ratio.'),
        ('p','The DFR-style comparator followed balanced last-layer refitting principles: frozen features, a source-only standardization transform, L1 logistic heads on equal tone-by-outcome group samples, and coefficient averaging over20 refits. C was selected using three inner folds and the mean worst-six-group accuracy, never the outer test data. This adapts DFR to MedSigLIP; it does not reproduce the original spurious-correlation benchmarks or train another encoder.'),
        ('p','E2 used20 nested outcome-stratified V-VI panels per budget of25,50 and100 images, followed by all available non-test V-VI images. E3 paired V-VI and lighter-tone panels by the exact released diagnosis and outcome. Each cell retained the smaller available count, giving identical diagnosis distributions and annotation budgets in each paired fit.'),
        ('p','Intervals used5,000 shared image bootstrap samples within outer-fold-by-tone-by-outcome strata, together with panel resampling for sampled-panel arms. Point estimates are means of panel-level AUROCs. Effects were calculated before rounding. These are unadjusted exploratory95% intervals conditional on the fitted heads and saved folds; they are not independent external replications.'),
        ('small','DFR source: Kirichenko, Izmailov and Wilson. Last Layer Re-Training is Sufficient for Robustness to Spurious Correlations. ICLR2023. https://arxiv.org/abs/2204.02937; author implementation linked in the recorded protocol.')]
    pages.append(p1)
    e1table=[['Method','DDI labels available*','V-VI AUROC [95% CI]','Delta vs PAD [95% CI]']]
    for a in e1:
        if a=='PAD_baseline':nb='0';de='Reference'
        else:
            if a in ['target_weighted','target_unweighted','lighter_weighted']:nb='165-166'
            elif a=='all_DDI_budget_matched':nb='165-166'
            else:nb='524-525'
            de=delta(a,'PAD_baseline') if ((cv.left_arm==a)&(cv.right_arm=='PAD_baseline')).any() else ''
        e1table.append([LABELS[a],nb,metric(a),de])
    p2=[('h','E1. Stronger baselines and annotation budget'),
        ('p',f'On the207 held-out V-VI images, including48 malignant cases, weighted targeted addition reached{metric("target_weighted")}. The DDI-only six-group-balanced head reached{metric("DDI_only_group_balanced")}. Its larger annotation budget must be stated explicitly; the results do not establish targeted addition as the best discrimination method.'),
        ('table',('Table E1. V-VI performance and annotation budget.',e1table,[68,22,42,42])),
        ('small','*Labels available within each non-test outer fold. DFR consumes the full labeled pool for internal tuning/group construction; each final component trains on an equal-size balanced subset. Full-DDI controls therefore have a larger available annotation budget. All use the same held-out V-VI images. Panel means and single-vector estimates are distinguished in the numerical CSV.'),
        ('p',f'Targeted addition exceeded the outcome- and budget-matched all-tone control by{delta("target_weighted","all_DDI_budget_matched")}. However, it did not clearly exceed full source-weighted all-DDI addition ({delta("target_weighted","all_DDI_source_weighted")}), tone-reweighted addition ({delta("target_weighted","all_DDI_group_reweighted")}), or the DFR-style refit ({delta("target_weighted","DFR_style")}). These comparisons support an annotation-efficiency interpretation rather than universal superiority.'),
        ('small','CIs for the inherited original score vectors were recomputed using the common extension bootstrap; original fitted scores and original archived tables were retained unchanged.')]
    pages.append(p2)
    e2table=[['V-VI training budget','Mean AUROC [95% CI]','Panel AUROC SD']]
    for label,a in [('0 (PAD-only)','PAD_baseline'),('25','target_budget_25'),('50','target_budget_50'),('100','target_budget_100'),('All:165-166','target_weighted')]:
        e2table.append([label,metric(a),f'{v.loc[a].AUROC_panel_SD:.3f}' if a.startswith('target_budget') else 'Single score vector'])
    p3=[('h','E2. Learning curve'),
        ('p',f'Mean V-VI AUROC increased from{v.loc["target_budget_25"].AUROC:.3f} at25 added images to{v.loc["target_budget_50"].AUROC:.3f} at50 and{v.loc["target_budget_100"].AUROC:.3f} at100. Using all165-166 available non-test images reached{v.loc["target_weighted"].AUROC:.3f}. The additional gain from100 to all available images was{delta("target_weighted","target_budget_100")}; a plateau was not established within this range.'),
        ('table',('Table E2. Nested annotation learning curve on held-out V-VI images.',e2table,[42,75,57])),
        ('image',('Figure_E2_learning_curve.png',148)),
        ('small','Figure E2. Mean performance over20 nested training panels; all-data value uses the saved full targeted model. Shading and bars show conditional image/panel95% CIs, not an external-cohort confidence band.'),
        ('h','E3. Exact diagnosis matching'),
        ('p',f'Exact diagnosis overlap retained114-125 images per source in each non-test fold ({coverage.target_fraction_retained.min()*100:.1f}-{coverage.target_fraction_retained.max()*100:.1f}% of the available V-VI adaptation images). The matched target and lighter arms achieved{metric("diagnosis_matched_target")} and{metric("diagnosis_matched_lighter")}, respectively. Their paired difference was{delta("diagnosis_matched_target","diagnosis_matched_lighter")}. This interval includes zero and meaningful positive effects; it does not establish a diagnosis-independent benefit.'),
        ('table',('Table E3. Diagnosis-matched adaptation on the same V-VI test images.',[
            ['Matched source','AUROC [95% CI]'],[LABELS['diagnosis_matched_target'],metric('diagnosis_matched_target')],
            [LABELS['diagnosis_matched_lighter'],metric('diagnosis_matched_lighter')]], [90,84])),
        ('small','Support is restricted to exact diagnosis-overlap adaptation cells. Full V-VI test images are retained. The original outcome-matched and new diagnosis-matched comparisons have different adaptation supports and budgets; they are not identical causal-effect estimands.')]
    pages.append(p3)
    old=pd.read_csv(out/'results/Existing_R5_refitting_robustness.csv');original=old[(old.stratum=='V-VI')&(old.metric=='AUROC')]
    # The original500-refit table is copied with source hash, not rerun or misrepresented as new baseline replication.
    existing=original[original.arm_or_contrast=='target_weighted_minus_PAD_baseline']
    if existing.empty:existing=original[original.arm_or_contrast.str.contains('target_weighted.*PAD_baseline',regex=True)]
    stability=[]
    if len(existing):
        r=existing.iloc[0];stability=[('p',f'The previously verified500-refit analysis of the original targeted-versus-PAD contrast remains available: original AUROC difference{r.original_point_estimate:.3f}, bootstrap mean{r.bootstrap_mean:.3f}, and refitting95% interval[{r.CI_low:.3f}, {r.CI_high:.3f}]. This supports stability of that original comparison, not independent validation of the new controls or external generalization.')]
    p4=[('h','Interpretation suitable for the Discussion'),
        ('p','The added controls change the interpretation of the original improvement. Targeted V-VI labeling provided better discrimination than an equally sized, outcome-matched random sample from all DDI tones. Nevertheless, a full-DDI balanced head achieved higher V-VI discrimination, and exact diagnosis matching substantially reduced the apparent target-versus-lighter advantage. The study should therefore distinguish annotation efficiency from absolute performance and avoid attributing the entire improvement to skin-tone coverage alone.'),
        ('p','The contribution supported by these results is a controlled comparison of domain exposure, annotation budget, feature reweighting and diagnostic composition under a shared frozen representation. The present evidence does not establish a new encoder architecture, a causal skin-tone effect, clinical usefulness, or superior performance on unseen institutions. The representation has PAD exposure during pretraining, and released DDI identifiers do not support a patient-independent evaluation claim.'),
        *stability,
        ('h','External validation remains the decisive missing experiment'),
        ('p','E4 must evaluate the prespecified frozen heads on actual MRA-MIDAS clinical-image features, with a reviewed malignant-versus-benign histopathology endpoint, intermediate diagnoses excluded, image-to-lesion aggregation, and patient-cluster uncertainty. Seven deployable heads and the evaluator are supplied. Input validation has been tested, but no MRA-MIDAS performance number has been generated. Dataset-specific usage terms and cross-dataset patient overlap still require verification.'),
        ('h','Integration into the designated MDPI documents'),
        ('p','Insert the verified Methods/Results/Discussion passages and Tables E1-E3 into MAKE_SkinCancer_20261006_MDPI.docx and its Supplementary_MDPI companion after those actual files are recovered. Align table/figure numbers and the DFR citation to the existing bibliography. The completed formatting checklist A should be retained. This addendum cannot establish that the unavailable documents have been edited or that E4 is complete.'),
        ('small','Historical Table S13 remains unverified. If 0.739 and 0.720 are rounded to the nearest three decimals from the same paired estimand, their unrounded difference is between 0.018 and 0.020; a displayed +0.021 cannot be explained by rounding alone. The original unrounded estimates and paired predictions are required. The archived new similarity-group analysis is a separate analysis and does not resolve the historical row.'),
        ('h','Author-review package'),
        ('p','The numerical package includes the recorded protocol, 1,261 saved heads, 97,744 image-level prediction records, sampling panels, inner tuning splits, unrounded effects and confidence intervals, and independent verification. Source raw images, feature caches and original metadata are excluded. The code uses the previously saved numerical archive identified by its SHA256.'),
        ('small','Rights context: DDI published research-use terms restrict redistribution of the source dataset. No Zenodo/public dataset or model release was performed. MRA-MIDAS primary source: https://aimi.stanford.edu/datasets/mra-midas-Multimodal-Image-Dataset-for-AI-based-Skin-Cancer. DDI terms: https://ddi-dataset.github.io/. All figures in this addendum derive from the saved experimental CSV files; no image-generation tool was used.')]
    pages.append(p4)
    pages.append([('h','Figures for manuscript integration'),('image',('Figure_E1_baseline_comparison.png',152)),
                  ('small','Figure E1. V-VI AUROC and exploratory95% CIs for prespecified controls. Full-DDI methods have a larger annotation budget than the targeted and budget-matched additions; see Table E1.'),
                  ('image',('Figure_E3_diagnosis_sensitivity.png',152)),
                  ('small','Figure E3. Target-minus-lighter AUROC contrasts before and after restricting adaptation to exact diagnosis common support. Supports and budgets differ. The diagnosis-matched interval includes zero; this comparison does not identify a causal effect of skin tone.')])
    # Keep each ordinary number/word separated in reader-facing text.
    import re
    def clean(s):
        protected={}
        def reserve(match):
            token='__RESERVED_'+chr(65+len(protected))+'__'
            protected[token]=match.group(0)
            return token
        s=re.sub(r'https?://\S+|a2e188c2…ca4ac4',reserve,s)
        # Keep E1/E2/E3 identifiers, hashes and source URLs intact.
        s=re.sub(r'([A-Za-z]{2,})(?=\d)',r'\1 ',s)
        s=re.sub(r'(?<=\d)(?=[A-Za-z])',' ',s)
        for token,value in protected.items():s=s.replace(token,value)
        return s
    for pp in pages:
        for j,(kind,value) in enumerate(pp):
            if kind in ['title','h','p','small']:pp[j]=(kind,clean(value))
    doc=Document();section=doc.sections[0];section.page_height=Mm(297);section.page_width=Mm(210)
    for key in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(section,key,Mm(18))
    for st in ['Normal','Title','Heading 1','Heading 2']:
        doc.styles[st].font.name='Times New Roman';doc.styles[st].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing=1.03;doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(17);doc.styles['Title'].font.color.rgb=RGBColor.from_string('1F5A99')
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].font.color.rgb=RGBColor.from_string('1F5A99')
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
    styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='BodyMake',fontName='Times-Roman',fontSize=10.3,leading=13.5,spaceAfter=7))
    styles.add(ParagraphStyle(name='SmallMake',fontName='Times-Roman',fontSize=8.1,leading=10.5,spaceAfter=6))
    styles.add(ParagraphStyle(name='HeadMake',fontName='Times-Bold',fontSize=12.5,leading=16,textColor=colors.HexColor('#1F5A99'),spaceAfter=7))
    styles.add(ParagraphStyle(name='TitleMake',fontName='Times-Bold',fontSize=17,leading=21,textColor=colors.HexColor('#1F5A99'),spaceAfter=8))
    styles.add(ParagraphStyle(name='CellMake',fontName='Times-Roman',fontSize=8.1,leading=10.2,alignment=TA_LEFT))
    story=[];md=[]
    for page_number,nodes in enumerate(pages):
        if page_number:doc.add_page_break();story.append(PageBreak());md.append('\n---\n')
        for kind,value in nodes:
            if kind in ['title','h','p','small']:
                p=doc.add_paragraph(value,style='Title' if kind=='title' else ('Heading 1' if kind=='h' else 'Normal'))
                if kind=='small':
                    for run in p.runs:run.font.size=Pt(8.5)
                story.append(Paragraph(html.escape(value),styles[{'title':'TitleMake','h':'HeadMake','p':'BodyMake','small':'SmallMake'}[kind]]))
                md.append(('## ' if kind=='h' else '# ' if kind=='title' else '')+value+'\n')
            elif kind=='table':
                caption,data,widths=value;doc.add_paragraph(caption,style='Normal');story.append(Paragraph(html.escape(caption),styles['SmallMake']))
                table=doc.add_table(rows=1,cols=len(data[0]));table.style='Light Shading Accent 1';table.autofit=False
                for j,label in enumerate(data[0]):table.rows[0].cells[j].text=label
                for rr in data[1:]:
                    cells=table.add_row().cells
                    for j,label in enumerate(rr):cells[j].text=str(label)
                for row in table.rows:
                    for j,cell in enumerate(row.cells):
                        cell.width=Mm(widths[j])
                        for pp in cell.paragraphs:
                            pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2)
                            for run in pp.runs:run.font.size=Pt(8.5)
                        trPr=row._tr.get_or_add_trPr();no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
                repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
                ptdata=[[Paragraph(html.escape(str(x)),styles['CellMake']) for x in rr] for rr in data]
                tt=Table(ptdata,colWidths=[w*72/25.4 for w in widths],repeatRows=1,hAlign='LEFT')
                tt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E7EEF6')),('BOX',(0,0),(-1,-1),.35,colors.HexColor('#B9C7D6')),
                                       ('INNERGRID',(0,0),(-1,-1),.25,colors.HexColor('#D5DEE8')),('VALIGN',(0,0),(-1,-1),'TOP'),
                                       ('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
                story.extend([tt,Spacer(1,8)]);md.append(caption+'\n\n'+'\n'.join(['| '+' | '.join(rr)+' |' for rr in [data[0],['---']*len(data[0])]+data[1:]])+'\n')
            elif kind=='image':
                name,width=value;path=figdir/name;doc.add_picture(str(path),width=Mm(width));pic=doc.paragraphs[-1];pic.alignment=WD_ALIGN_PARAGRAPH.CENTER
                with Image.open(path) as im:w,h=im.size
                story.append(PDFImage(str(path),width=width*72/25.4,height=width*72/25.4*h/w));story.append(Spacer(1,5))
                md.append(f'Figure asset: figures/{name}\n')
    for pp in section.footer.paragraphs:
        pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;run=pp.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
    title='MAKE_Results_Addendum_20261006';doc.core_properties.title='Additional controlled experiments for the MAKE manuscript'
    doc.core_properties.subject='Exploratory E1-E3 results; E4 not evaluated; does not replace the designated MDPI manuscript'
    doc.save(out/'documents'/f'{title}.docx')
    def footer(canvas,doc_):
        canvas.saveState();canvas.setFont('Times-Roman',8);canvas.setFillColor(colors.HexColor('#617184'))
        canvas.drawString(51,25,'MAKE | Additional controlled experiments | 6 October 2026');canvas.drawRightString(A4[0]-51,25,str(doc_.page));canvas.restoreState()
    pdf=SimpleDocTemplate(str(out/'documents'/f'{title}.pdf'),pagesize=A4,leftMargin=18*72/25.4,rightMargin=18*72/25.4,
                          topMargin=18*72/25.4,bottomMargin=18*72/25.4,title=doc.core_properties.title,author='Author-review research package')
    pdf.build(story,onFirstPage=footer,onLaterPages=footer)
    (out/'documents'/f'{title}.md').write_text('\n'.join(md)+'\n')
    print('Figures saved at600dpi with CSV provenance; DOCX/PDF/Markdown insertion text generated.')


if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--output',required=True);a=ap.parse_args();build(a.output)
