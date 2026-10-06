"""Create the author-facing results report directly from saved analysis tables."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import make_reanalysis_common as c


def number(value, signed=False, digits=3):
    return format(float(value), ('+' if signed else '') + f'.{digits}f')


def bound(value):
    value = float(value)
    return format(value, '.5f' if 0 < abs(value) < 0.0005 else '.3f')


def ci(point, low, high, signed=False):
    return number(point, signed) + '\n[' + bound(low) + ', ' + bound(high) + ']'


def build(output_root, document_path):
    out = Path(output_root)
    plan = json.loads((out / 'plan/MAKE_Analysis_Plan_v1.json').read_text())
    arms = c.read_csv(out / 'results/Table_R3_DDI_arms.csv')
    controls = c.read_csv(out / 'results/Table_R4_matched_controls.csv')
    refits = c.read_csv(out / 'results/Table_R5_refitting_robustness.csv')
    draws = c.read_csv(out / 'results/matched_draw_summary.csv')
    proxy = c.read_csv(out / 'results/Table_RS13_proxy_group_changes.csv')
    proxy_arms = c.read_csv(out / 'results/Table_RS13_proxy_group_arms.csv')
    proof = json.loads((out / 'logs/independent_verification.json').read_text())
    primary = next(r for r in controls if r['stratum']=='V-VI' and r['metric']=='AUROC' and r['left_arm']=='target_weighted' and r['right_arm']=='PAD_baseline')
    retrain = next(r for r in refits if r['stratum']=='V-VI' and r['metric']=='AUROC' and r['arm_or_contrast']=='target_weighted_minus_PAD_baseline')
    advantage = next(r for r in controls if r['stratum']=='V-VI' and r['metric']=='AUROC' and r['left_arm']=='target_weighted' and r['right_arm']=='lighter_weighted')
    secondary = json.loads((out / 'plan/MAKE_Proxy_Group_Plan_v1.json').read_text())

    d = Document()
    section = d.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.9)
    normal = d.styles['Normal']
    normal.font.name, normal.font.size = 'Times New Roman', Pt(11)
    normal.font.color.rgb = RGBColor(0,0,0)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(7)
    for style_name, size in [('Title',20),('Heading 1',15),('Heading 2',12),('Caption',9.5)]:
        style=d.styles[style_name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(8)
        style.paragraph_format.keep_with_next=True
    d.styles['Title'].font.bold=True
    d.styles['Caption'].paragraph_format.keep_with_next=False
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run();run.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
    d.core_properties.title='Skin lesion reanalysis for MAKE'
    d.core_properties.subject='Numerical controlled comparisons and traceable new analysis'
    d.core_properties.author='Sakchai Tangprasert'

    def paragraph(text, bold=False):
        p=d.add_paragraph();p.add_run(text).bold=bold;return p

    start_new_page = False

    def heading(text):
        nonlocal start_new_page
        p = d.add_heading(text,level=1)
        if start_new_page:
            p.paragraph_format.page_break_before = True
            start_new_page = False

    def page():
        nonlocal start_new_page
        start_new_page = True

    def table(headers, rows, widths):
        t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for i,w in enumerate(widths):t.columns[i].width=Inches(w)
        for i,header in enumerate(headers):t.rows[0].cells[i].text=header
        trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trpr.append(repeat)
        for row in rows:
            cells=t.add_row().cells
            for i,value in enumerate(row):cells[i].text=str(value)
        borders=OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        t._tbl.tblPr.append(borders)
        for ri,row in enumerate(t.rows):
            for col,cell in enumerate(row.cells):
                cell.width=Inches(widths[col]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
                for edge,size in [('top',75),('bottom',75),('left',85),('right',85)]:
                    e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(size));e.set(qn('w:type'),'dxa');margins.append(e)
                props.append(margins)
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'DFEAF2' if ri==0 else ('F8F8F8' if ri%2==0 else 'FFFFFF'));props.append(fill)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=False
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if col==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:run.font.size=Pt(10.5);run.font.bold=ri==0;run.font.color.rgb=RGBColor(0,0,0)
        d.add_paragraph().paragraph_format.space_after=Pt(0)
        return t

    def figure(stem, caption, width=6.6):
        p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(out/'figures'/(stem+'.png')),width=Inches(width))
        d.add_paragraph(caption,style='Caption')

    d.add_paragraph('Skin lesion reanalysis for MAKE',style='Title')
    paragraph('5 October 2026',bold=False)
    paragraph(f'In this new exploratory analysis, adding weighted Fitzpatrick V-VI DDI training images increased V-VI AUROC from {number(primary["right_value"])} to {number(primary["left_value"])}. The paired change was {number(primary["difference_unrounded"],True)}, with a conditional 95% interval of {bound(primary["CI_low"])} to {bound(primary["CI_high"])}. Refitting the heads in 500 bootstrap replicates gave a 95% interval of {bound(retrain["CI_low"])} to {bound(retrain["CI_high"])} for the same original point estimate.')
    paragraph('These results use the full curated PAD pool and newly generated image features. DDI is a previously studied cohort, and its released metadata lack patient identifiers. The findings describe within-DDI adaptation and cannot establish clinical referral performance or independence between DDI patients.')
    heading('Cohorts and transfer')
    table(['Dataset and skin tone','Images','Malignant label','Benign label'],[
          ['PAD curated pool',2268,1070,1198],['DDI I-II',208,49,159],['DDI III-IV',241,74,167],['DDI V-VI',207,48,159]], [2.7,1.0,1.5,1.5])
    paragraph('PAD contains 1,370 source patient groups. Its 804 images with missing skin-tone labels all have negative released labels and were retained without imputation. Only 11 PAD images are labelled V-VI. MedSigLIP used PAD during pretraining [3], which limits interpretation of PAD development performance as evidence of transfer.')
    figure('Figure_R1_transfer','Figure R1. Conditional bootstrap intervals for PAD nested development and PAD-only DDI transfer. PAD uses patient-cluster resampling; DDI uses image resampling stratified by fold, tone and outcome. Each interval uses 10,000 resamples.')

    page();heading('Controlled comparison at Fitzpatrick V to VI')
    paragraph('All five arm summaries below evaluate the same 207 V-VI images, including 48 with positive released labels. Added rows come from outside each saved test fold. Each lighter panel exactly matches the number of positive and negative V-VI training rows. C=10 is selected from PAD only and held fixed across these arms.')
    labels={'PAD_baseline':'PAD only','target_unweighted':'Add V-VI without source weight','target_weighted':'Add V-VI with source weight','lighter_unweighted':'Add lighter tones without source weight\nMean of 20 draws','lighter_weighted':'Add lighter tones with source weight\nMean of 20 draws'}
    rows=[]
    for name in ['PAD_baseline','target_unweighted','target_weighted','lighter_unweighted','lighter_weighted']:
        a=next(r for r in arms if r['stratum']=='V-VI' and r['arm']==name)
        delta='Reference'
        if name!='PAD_baseline':
            z=next(r for r in controls if r['stratum']=='V-VI' and r['metric']=='AUROC' and r['left_arm']==name and r['right_arm']=='PAD_baseline')
            delta=ci(z['difference_unrounded'],z['CI_low'],z['CI_high'],True)
        rows.append([labels[name],ci(a['AUROC'],a['AUROC_CI_low'],a['AUROC_CI_high']),delta,number(100*float(a['sensitivity_FPR_le_0_50']),digits=1)])
    table(['Training arm','AUROC\n[95% CI]','Change from PAD\n[95% CI]','Sensitivity at FPR <=0.50 (%)'],rows,[2.25,1.4,1.55,1.5])
    paragraph(f'The weighted target-minus-lighter contrast was {number(advantage["difference_unrounded"],True)} [{bound(advantage["CI_low"])}, {bound(advantage["CI_high"])}]. Arm intervals are conditional on fitted heads; changes use paired resamples. They exclude head-refitting uncertainty and use no multiplicity correction.')
    paragraph('Sensitivity uses a retrospective ROC threshold chosen on labelled evaluation scores to keep empirical false positives at or below 50% of negatives. The achieved V-VI FPR was 79/159, or 49.7%. This operating point is not an independently validated deployment threshold.')
    page();heading('Controlled contrasts and repeated draws')
    figure('Figure_R2_controls','Figure R2. V-VI AUROC contrasts with 10,000 paired conditional bootstrap intervals. The first four rows compare additions with the PAD-only head. Other rows compare the two named additions. Lighter arms average per-draw metrics, rather than scores.',width=6.6)
    figure('Figure_R4_repeated_draws','Figure R4. V-VI AUROC changes across 20 matched lighter-panel draws, with means marked. All draws use the same held-out DDI images. The dots are not independent evaluation cohorts.',width=6.45)

    page();heading('Sampling variability and head refitting')
    paragraph('The 20 lighter panels vary training sample selection while retaining the same evaluation images. Their variation is described separately from a confidence interval over evaluation images.')
    rows=[]
    for arm in ['lighter_unweighted','lighter_weighted']:
        a=next(r for r in draws if r['stratum']=='V-VI' and r['arm']==arm)
        rows.append(['Lighter without source weight' if arm.endswith('unweighted') else 'Lighter with source weight',
                    number(a['AUROC_mean']),number(a['AUROC_SD'],digits=4),number(a['AUROC_minimum'])+' to '+number(a['AUROC_maximum']),number(a['AUROC_mean_change_from_PAD_baseline'],True)])
    table(['Matched lighter panels','Mean AUROC','SD over 20 draws','AUROC range','Mean change'],rows,[2.25,1.05,1.1,1.25,1.05])
    paragraph('A second analysis refitted one PAD-only head and five weighted V-VI adaptation heads in each of 500 paired bootstrap replicates. PAD patients were sampled as clusters; DDI images were sampled within original fold, tone and outcome strata. The encoder, C and fold assignments remained fixed.')
    rows=[]
    for name,label in [('PAD_baseline','PAD only'),('target_weighted','Weighted V-VI addition'),('target_weighted_minus_PAD_baseline','Paired AUROC change')]:
        a=next(r for r in refits if r['stratum']=='V-VI' and r['metric']=='AUROC' and r['arm_or_contrast']==name)
        rows.append([label,number(a['original_point_estimate'],name.endswith('baseline') and name!='PAD_baseline'),
                    number(a['bootstrap_mean'],name.endswith('baseline') and name!='PAD_baseline'),bound(a['CI_low'])+' to '+bound(a['CI_high'])])
    table(['Refitting estimate','Original point','Bootstrap mean','Refitting 95% CI'],rows,[2.25,1.15,1.15,2.15])
    paragraph('The original paired point estimate and the mean over refitted replicates are distinct summaries. The intervals include the specified head-refitting and sampling steps, with a finite 500-replicate approximation. They exclude C selection, fold selection, encoder uncertainty and unobserved patient dependence within DDI. This is a robustness analysis within existing data.')
    figure('Figure_R3_uncertainty','Figure R3. Uncertainty for the original weighted V-VI AUROC change. Fixed-head, head-refitting and post hoc similarity-group analyses are labelled separately. No row represents independent external replication.',width=6.6)

    page();heading('Image similarity and arithmetic')
    paragraph(f'The raw-image DCT-hash screen flagged four near-identical candidate pairs. Three pairs crossed original folds. We declared a separate post hoc sensitivity plan, joined each connected component, and assigned it to its minimum original fold. All singleton assignments stayed unchanged. This created {secondary["image_similarity_components"]} image-similarity components and moved {secondary["images_moved_to_another_fold"]} images between folds.')
    paragraph('One candidate pair, DDI:000297.png and DDI:000546.png, has different released tone groups, I-II and III-IV. The original annotations were retained. All four pairs have matching released malignant labels. These findings do not establish patient identities, and grouping similar photographs cannot exclude less-similar photographs of the same patient.')
    rows=[]
    for tone in ['I-II','III-IV','V-VI']:
        a=next(r for r in proxy if r['stratum']==tone and r['arm']=='target_weighted' and r['metric']=='AUROC')
        info=next(r for r in proxy_arms if r['stratum']==tone and r['arm']=='target_weighted')
        rows.append([tone,info['n_proxy_components'],number(a['baseline'],digits=6),number(a['adapted'],digits=6),
                    number(a['difference_unrounded'],True,digits=6),bound(a['CI_low'])+' to '+bound(a['CI_high'])])
    table(['Tone','Proxy components','PAD AUROC','After V-VI addition','Paired change','Conditional 95% CI'],rows,[0.55,0.9,1.1,1.25,1.15,1.75])
    paragraph('Table RS13. New post hoc similarity-group results for weighted V-VI addition evaluated across tones. Changes are computed before display rounding; the CSV retains full precision. One component spans I-II and III-IV, so per-tone component counts are not additive. The V-VI point estimate stays unchanged because no flagged pair involves V-VI. The intervals use 10,000 shared component bootstrap samples and remain conditional on fitted heads.')
    heading('Historical Table S13')
    paragraph('The reviewed supplement reports 0.720 before and 0.739 after the I-II intervention, with a change of +0.021. Subtracting the displayed values gives +0.019. Ordinary three-decimal rounding of the same point estimator permits an underlying difference approximately between +0.018 and +0.020. It cannot explain +0.021 by itself. The unrounded historical outputs are unavailable, so +0.019 is a displayed-value calculation, not a verified replacement estimate.')
    paragraph('The new Table RS13 above evaluates V-VI addition in a different curated training pool and with newly specified image-similarity groups. The historical Table S13 instead describes stratum-specific additions and different proxy groups. Its numerical result and interval remain unverified. The new evidence should be reported with its own caption and methods.')

    page();heading('Methods for the new analysis')
    paragraph('We used 2,268 curated PAD-UFES-20 images [1] after quarantining all 30 images in exact-duplicate groups, and all 656 released DDI images [2]. PAD BCC, SCC and MEL were labelled positive; ACK, SEK and NEV were labelled negative. DDI malignant Boolean values were preserved. These are dataset-defined malignancy outcomes rather than tested clinical referral urgency.')
    paragraph('The frozen encoder was google/medsiglip-448 at revision 9cea28a1a1195f665105faa6e8544c112fd960a4 [3]. PIL decoded each image and converted it to RGB, discarding alpha, without EXIF transposition or augmentation. The recorded PIL processor resized and normalised images at 448 by 448. Image-tower pooler outputs were L2-normalised into 1,152 float32 features, then cast to float64 for head fitting. Model and source-image hashes, processor configuration and the actual CUDA environment are retained in the extraction audit.')
    paragraph('L2 logistic regression used lbfgs, l1_ratio=0, an intercept and tolerance 1e-7 [4]. The iteration budget started at 4,000, with same-objective retries at larger budgets only if needed. We selected C over 0.001, 0.01, 0.1, 1, 10, 100 and 1000 by mean validation AUROC on five PAD patient-grouped folds. Nested PAD development additionally selected C inside each outer training fold. Final C=10 came from full-PAD development folds and was fixed before DDI head fitting.')
    paragraph('For each DDI fold, the adaptation head used all PAD images plus V-VI DDI rows outside that fold. Unweighted addition used source factor 1. Weighted addition used sqrt(N_PAD/N_added). Explicit outcome-class factors were N/(2*n_class), based on unweighted training-row counts, before multiplying source factors. The same explicit formula, updated for bootstrap multiplicities, was used in refitting. Weighted versus unweighted arms also change total loss weight and effective penalty. Matched target-versus-lighter arms use equal row counts, outcome counts and source multipliers.')
    paragraph('The control used 20 independently seeded draws without replacement from I-II and III-IV DDI training rows. Each draw matched the number of positive and negative V-VI added rows in that fold. All DDI evaluation-fold rows were excluded from every corresponding fit. Reported lighter-arm metrics are the mean of per-draw AUROCs and sensitivities. Averaging their scores would define a different model and was not used.')
    paragraph('The main uncertainty analysis used 10,000 shared DDI image bootstrap counts stratified by fold, tone and label. PAD development used 10,000 patient-cluster bootstrap counts. The 500-replicate refitting analysis preserved original folds and C while refitting six heads per replicate. All intervals are percentile 95% intervals without multiple-comparison adjustment. Sensitivity is read retrospectively at the most permissive attainable empirical threshold with FPR at or below 0.50, excluding ties at its boundary. DDI thresholds were not tuned as deployment thresholds.')
    paragraph('The main plan was hash-locked before this run fitted or inspected new classifier scores. It is an analysis record, not a registered preregistration. The image-similarity sensitivity was declared after the primary results were inspected and is reported as post hoc. Unknown DDI patients and possible unobserved image overlap remain unresolved.')

    page();heading('Traceability and manuscript changes')
    paragraph(f'The verification reconstructed {proof["saved_heldout_and_tuning_logits_reconstructed"]:,} saved held-out and tuning logits from {proof["classifier_heads_with_coefficients_and_fit_weights_verified"]:,} classifier heads. It checked image and feature order, coefficients, training-row multiplicities, effective weights, held-out separation, matched outcome counts and reported unrounded contrasts. Bootstrap metrics were also checked against literal resampling and an independent ROC implementation. The numerical verification completed with status PASS.')
    paragraph('The accompanying ZIP contains fixed feature matrices, full source-preserving indices, saved folds, selected lighter panels, classifier coefficient files, per-image logits, uncertainty samples, exact scripts and audit logs. Figures R1-R4 have plotted data CSV files and machine-readable source and output hashes. The code regenerates the charts from numerical results.')
    paragraph('The 3.51 GB source encoder weights are referenced by their exact model revision and SHA256 rather than included in the ZIP. The source configurations were independently fetched and matched the extraction audit. A Drive folder listing did not expose the weight file, so remote persistence of that raw checkpoint was not independently confirmed. Numerical verification here starts from the bundled saved feature matrices. The supplied loader can retrieve the recorded model revision under the existing authorised account.')
    paragraph('Table 4 can be revised using the new controlled comparison above, provided the Methods and all related statements specify the full curated PAD pool and the new DDI protocol. New numbers should not be combined with the historical combined-pool AUROC, training subsets or unverified cross-encoder claims.')
    paragraph('Historical Table 5 describes a held-out Fitzpatrick17k subset of 438 images, including 311 positives and 424 proxy clusters. Its raw subset manifest, predictions and fitted model were not available in the supplied material. Numerical replication results and confidence intervals cannot be reconstructed from those counts. This report therefore provides no verified external-replication result for that table. The 500-replicate robustness analysis must retain its own label.')
    paragraph('Historical Table S13 needs the original unrounded before-and-after estimates and paired score files to resolve its arithmetic discrepancy. The new grouped sensitivity can be cited separately with its own group definition. Other historical numerical claims and graph provenance are not verified by this new run.')
    heading('Primary sources')
    for text in ['[1] PAD-UFES-20 release version 1. https://data.mendeley.com/datasets/zr7vgbcyr2/1',
                 '[2] Diverse Dermatology Images dataset. https://ddi-dataset.github.io/',
                 '[3] Google MedSigLIP model card. https://huggingface.co/google/medsiglip-448',
                 '[4] scikit-learn 1.8 LogisticRegression documentation. https://scikit-learn.org/1.8/modules/generated/sklearn.linear_model.LogisticRegression.html']:
        p=paragraph(text);p.paragraph_format.space_after=Pt(4)
        for run in p.runs:run.font.size=Pt(9.5)
    path=Path(document_path);path.parent.mkdir(parents=True,exist_ok=True);d.save(path)
    c.write_file(out/'code'/Path(__file__).name,Path(__file__).read_bytes())
    print('REPORT CREATED',path.resolve())


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',required=True);p.add_argument('--document',required=True)
    args=p.parse_args();build(args.output,args.document)
