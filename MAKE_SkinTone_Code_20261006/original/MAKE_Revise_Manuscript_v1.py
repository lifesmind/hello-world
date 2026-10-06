"""Revise the MAKE manuscript and supplement from verified reanalysis outputs.

No fitting or statistical resampling is performed here. Every numerical table is
constructed from saved source CSV files. Original author declarations are retained
through the source profile; editorial explanations are native Word comments.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE = 'Controlled Data Addition for Skin Lesion Classification Across Skin Tones Using Frozen MedSigLIP Features'
ARM_LABELS = {
    'PAD_baseline': 'PAD only',
    'target_unweighted': 'V-VI addition without source weight',
    'target_weighted': 'V-VI addition with source weight',
    'lighter_unweighted': 'Lighter addition without source weight',
    'lighter_weighted': 'Lighter addition with source weight',
}
ARM_ORDER = list(ARM_LABELS)
TONES = ['I-II', 'III-IV', 'V-VI']
TONE_ORDER = ['ALL_DDI'] + TONES


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_csv(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def n(x, signed=False, digits=3):
    return format(float(x), ('+' if signed else '') + f'.{digits}f')


def b(x):
    x = float(x)
    return format(x, '.5f' if 0 < abs(x) < 0.0005 else '.3f')


def interval(point, low, high, signed=False):
    return n(point, signed) + '\n(' + b(low) + ' to ' + b(high) + ')'


def rgb_font(run, size=11, bold=None):
    run.font.name = 'Palatino Linotype'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), 'Palatino Linotype')
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'csTheme'):
        fonts.attrib.pop(qn('w:' + key), None)


class Article:
    def __init__(self, refs, citation_order, supplement=False):
        self.d = Document()
        self.refs, self.order = refs, citation_order
        self.tables, self.comments = [], []
        sec = self.d.sections[0]
        # Preserve the source manuscript's A4 format and Palatino typography.
        sec.page_width, sec.page_height = Inches(8.2681), Inches(11.6931)
        sec.top_margin, sec.bottom_margin = Inches(.78), Inches(.7)
        sec.left_margin = sec.right_margin = Inches(.7)
        sec.header_distance, sec.footer_distance = Inches(.28), Inches(.3)
        for name, size in [('Normal', 11), ('Title', 17), ('Heading 1', 12.5), ('Heading 2', 11.5), ('Caption', 9.5)]:
            style = self.d.styles[name]
            style.font.name, style.font.size = 'Palatino Linotype', Pt(size)
            style.font.color.rgb = RGBColor(0, 0, 0)
            pf = style.paragraph_format
            pf.line_spacing, pf.space_after = 1.08, Pt(6)
            pf.left_indent = pf.right_indent = Inches(0)
            pf.widow_control = True
            if name != 'Normal':
                pf.keep_with_next = True
            if name.startswith('Heading'):
                style.font.bold = True; pf.space_before = Pt(11)
            if name == 'Caption':
                pf.keep_with_next = False
            if style.element.pPr is not None:
                for tag in ('pBdr', 'shd'):
                    for el in style.element.pPr.findall(qn('w:' + tag)):
                        style.element.pPr.remove(el)
        self.d.styles['Title'].font.bold = True
        self.d.core_properties.title = TITLE + (' Supplementary Materials' if supplement else '')
        self.d.core_properties.author = 'Sakchai Tangprasert and Suwit Chanthasen'
        hp = sec.header.paragraphs[0]
        hp.text = 'Machine Learning and Knowledge Extraction' + ('  Supplementary Materials' if supplement else '')
        for run in hp.runs: rgb_font(run, 8)
        footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

    def cite(self, text):
        def sub(m):
            keys = m.group(1).split(',')
            values = []
            for key in keys:
                assert key in self.refs, key
                if key not in self.order: self.order.append(key)
                values.append(str(self.order.index(key) + 1))
            return '[' + ','.join(sorted(values, key=int)) + ']'
        return re.sub(r'\{\{([a-z,]+)\}\}', sub, text)

    def p(self, text, size=11, bold=False, style=None, keep=False):
        p = self.d.add_paragraph(style=style)
        if style is None: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.keep_with_next = keep
        rgb_font(p.add_run(self.cite(text)), size, bold)
        return p

    def h(self, text, level=1, new_page=False):
        p = self.d.add_heading(text, level=level)
        for run in p.runs: rgb_font(run, 12.5 if level == 1 else 11.5, True)
        p.paragraph_format.page_break_before = new_page
        return p

    def comment(self, paragraph, text):
        comment = self.d.add_comment(paragraph.runs, text=text, author='Editorial review', initials='ER')
        self.comments.append({'comment_id': comment.comment_id, 'anchor': paragraph.text, 'text': text})

    def table(self, label, caption, headers, rows, widths, source):
        cap = self.p(label + '. ' + caption, size=9.5, bold=True, keep=True)
        cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        t = self.d.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        for i, width in enumerate(widths): t.columns[i].width = Inches(width)
        for i, text in enumerate(headers): t.rows[0].cells[i].text = text
        for row in rows:
            for cell, text in zip(t.add_row().cells, row, strict=True): cell.text = str(text)
        repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement('w:' + edge); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9'); borders.append(el)
        t._tbl.tblPr.append(borders)
        group_size = {'Table S7': 5, 'Table S8': 5, 'Table S9': 4, 'Table S11': 3}.get(label)
        for ri, row in enumerate(t.rows):
            # Keep each record intact; long tables may break between records.
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for j, cell in enumerate(row.cells):
                cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp = cell._tc.get_or_add_tcPr(); margins = OxmlElement('w:tcMar')
                for edge, value in [('top', 75), ('bottom', 75), ('left', 85), ('right', 85)]:
                    el = OxmlElement('w:' + edge); el.set(qn('w:w'), str(value)); el.set(qn('w:type'), 'dxa'); margins.append(el)
                cp.append(margins)
                fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'E8EDF2' if ri == 0 else ('F8F8F8' if ri % 2 == 0 else 'FFFFFF')); cp.append(fill)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.keep_with_next = ri < len(t.rows)-1 and (len(rows) <= 8 or ri == 0 or (group_size is not None and ri % group_size != 0))
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 or label == 'Table S1' else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs: rgb_font(run, 9.5, ri == 0)
        self.p('', size=3).paragraph_format.space_after = Pt(0)
        self.tables.append({'label': label, 'source': source, 'headers': headers, 'rows': rows})
        return cap

    def figure(self, label, file, caption, width=6.5):
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(3)
        p.add_run().add_picture(str(file), width=Inches(width))
        pic = p._p.xpath('.//wp:docPr')[0]
        pic.set('descr', label + '. ' + caption)
        self.p(label + '. ' + caption, size=9.5, style='Caption')


def draw_flow(root):
    fig, ax = plt.subplots(figsize=(9.2, 4.0)); ax.set(xlim=(0, 10), ylim=(0, 6)); ax.axis('off')
    boxes = [
        (2.55, 4.8, 4.9, .82, 'Frozen MedSigLIP image features\nPAD 2268 images and DDI 656 images'),
        (.25, 3.4, 4.55, .86, 'PAD patient grouped development\nNested evaluation and final C selection'),
        (5.2, 3.4, 4.55, .86, 'DDI five saved image folds\nStratified by skin tone and outcome'),
        (.25, 1.93, 4.55, .92, 'PAD only baseline\nFit all PAD and score all DDI'),
        (5.2, 1.93, 4.55, .92, 'Adaptation and matched controls\nAll PAD plus other DDI folds only'),
        (2.55, .30, 4.9, .95, 'Held out DDI scores and paired contrasts\nConditional bootstrap and head refitting'),
    ]
    for x, y, w, h, text in boxes:
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0.03,rounding_size=.06', facecolor='#F5F7F8', edgecolor='#687680', linewidth=1))
        ax.text(x+w/2, y+h/2, text, ha='center', va='center', fontsize=10, color='black')
    for start, end in [((3.8,4.8),(2.5,4.28)),((6.2,4.8),(7.5,4.28)),((2.5,3.4),(2.5,2.89)),((7.5,3.4),(7.5,2.89)),((4.8,3.7),(5.2,2.75)),((2.5,1.93),(3.8,1.28)),((7.5,1.93),(6.2,1.28))]:
        ax.annotate('', xy=end, xytext=start, arrowprops={'arrowstyle':'->','color':'#687680','lw':1.2})
    fig.tight_layout(pad=.5)
    folder=root/'figures'; folder.mkdir(exist_ok=True)
    for suffix in ('png','svg'): fig.savefig(folder/('Figure_M1_study_flow.'+suffix), dpi=300, bbox_inches='tight')
    plt.close(fig)
    provenance={'figure':'Figure_M1_study_flow','type':'original schematic of the executed analysis',
        'authoring_code_sha256':sha(__file__),'plan_sha256':sha(root/'plan/MAKE_Analysis_Plan_v1.json'),
        'source_counts_sha256':sha(root/'inputs/preparation_counts/dataset_counts.csv'),
        'outputs_sha256':{suffix:sha(folder/('Figure_M1_study_flow.'+suffix)) for suffix in ('png','svg')}}
    (folder/'study_flow_provenance.json').write_text(json.dumps(provenance,indent=2)+'\n')


def build(root, main_path, supp_path):
    root=Path(root)
    profile=json.loads((root/'inputs/paper_source_profile.json').read_text())
    assert json.loads((root/'logs/independent_verification.json').read_text())['status']=='PASS'
    data={name:read_csv(root/'results'/name) for name in [
        'Table_R2_PAD_development_DDI_transfer.csv','Table_R3_DDI_arms.csv','Table_R4_matched_controls.csv',
        'Table_R5_refitting_robustness.csv','Table_R6_repeated_draws.csv','matched_draw_summary.csv',
        'Table_RS13_proxy_group_changes.csv','Table_RS13_proxy_group_arms.csv','PAD_tuning_fold_AUROC.csv',
        'DDI_similarity_candidate_pairs.csv','DDI_similarity_metadata_consistency.csv']}
    counts=read_csv(root/'inputs/preparation_counts/dataset_counts.csv')
    folds=read_csv(root/'inputs/preparation_counts/fold_counts.csv')
    fit=json.loads((root/'logs/classifier_fit_audit.json').read_text())
    verification=json.loads((root/'logs/independent_verification.json').read_text())
    spec=json.loads((root/'inputs/features/feature_spec.json').read_text())
    plan=json.loads((root/'plan/MAKE_Analysis_Plan_v1.json').read_text())
    secondary=json.loads((root/'plan/MAKE_Proxy_Group_Plan_v1.json').read_text())
    def arm(tone, name): return next(r for r in data['Table_R3_DDI_arms.csv'] if r['stratum']==tone and r['arm']==name)
    def contrast(tone, left, right, metric='AUROC'): return next(r for r in data['Table_R4_matched_controls.csv'] if r['stratum']==tone and r['left_arm']==left and r['right_arm']==right and r['metric']==metric)
    def refit(tone, name, metric='AUROC'): return next(r for r in data['Table_R5_refitting_robustness.csv'] if r['stratum']==tone and r['arm_or_contrast']==name and r['metric']==metric)
    primary=contrast('V-VI','target_weighted','PAD_baseline')
    target_lighter=contrast('V-VI','target_weighted','lighter_weighted')
    uncertainty=refit('V-VI','target_weighted_minus_PAD_baseline')
    contrast_pairs=[('target_weighted','PAD_baseline'),('target_unweighted','PAD_baseline'),('lighter_weighted','PAD_baseline'),('lighter_unweighted','PAD_baseline'),('target_weighted','lighter_weighted'),('target_unweighted','lighter_unweighted'),('target_weighted','target_unweighted')]
    draw_flow(root)
    order=[];m=Article(profile['references'],order)
    m.p('Article',size=10,bold=True)
    title=m.p(TITLE,size=17,bold=True,style='Title',keep=True)
    m.comment(title,'The title now describes the verified binary malignancy classification experiment. Historical referral-triage, cross-encoder and seven-source claims were removed because the current reanalysis cannot verify them.')
    for text in profile['frontmatter']: m.p(text,size=9.5)
    abstract=(f'Abstract: We evaluated whether adding clinical images from a target skin-tone stratum improves binary skin-lesion classification beyond matched same-source additions. Frozen MedSigLIP features were extracted for 2,268 curated PAD-UFES-20 images and 656 Diverse Dermatology Images (DDI). An L2 logistic-regression head used C selected only from patient-grouped PAD development. Within five saved DDI image folds, we compared a PAD-only baseline, Fitzpatrick V-VI additions with and without source weighting, and 20 lighter-tone panels matched for positive and negative training counts. Baseline V-VI AUROC was {n(primary["right_value"])} and increased to {n(primary["left_value"])} with weighted V-VI addition. The paired change was {n(primary["difference_unrounded"],True)} (conditional 95% confidence interval, {b(primary["CI_low"])} to {b(primary["CI_high"])}); its interval after head refitting in 500 bootstrap replicates was {b(uncertainty["CI_low"])} to {b(uncertainty["CI_high"])}. The weighted target-minus-lighter contrast was {n(target_lighter["difference_unrounded"],True)} ({b(target_lighter["CI_low"])} to {b(target_lighter["CI_high"])}). A post hoc similarity-group analysis retained the direction of the V-VI change. These results describe exploratory adaptation within DDI. Unknown patient overlap, encoder pretraining exposure and differences in disease composition prevent a causal skin-tone interpretation or an independent clinical-validation claim.')
    ap=m.p(abstract,size=10.5)
    m.comment(ap,'All abstract estimates now come from the saved new analysis. The 500-refit interval concerns the original paired point estimate, not the mean of the refitted bootstrap estimates. No external Fitzpatrick17k replication is claimed.')
    m.p('Keywords: skin lesion classification; dataset shift; skin tone; MedSigLIP; matched data addition; bootstrap uncertainty',size=10)
    m.h('1 Introduction')
    m.p('Skin-lesion classifiers can perform differently across image sources and patient groups even when their development accuracy is high. Public dermatology datasets often document skin tone incompletely, which limits subgroup evaluation {{wen}}. The Diverse Dermatology Images (DDI) study established that strong dermatology models can lose discrimination on a clinically curated and diverse benchmark, with poorer results on darker skin and uncommon conditions {{ddi}}. Clinician decision-support research likewise shows that performance and the effect of assistance depend on skin tone and the evaluation setting {{groh}}. These findings motivate testing which changes to training data improve a fixed classification pipeline.')
    m.p('A target-domain addition may help because it represents the target subgroup, because it supplies images from the same source, or because it changes the training outcome distribution. A comparison with no added data cannot separate these possibilities. We therefore compare V-VI clinical-image additions with lighter-tone DDI additions matched for the number of positive and negative rows. We also hold the encoder and PAD-selected penalty parameter fixed, and apply the same source-weighting rule to the matched additions.')
    m.p('We use smartphone clinical images from PAD-UFES-20 to develop the classifier {{pad,padrelease}} and the released DDI cohort for transfer evaluation and within-cohort adaptation. MedSigLIP supplies a fixed medical image representation {{modelcard}}. The questions are whether PAD-only discrimination varies across DDI tone strata, whether V-VI addition improves V-VI discrimination more than a matched lighter-tone addition, and how the estimated gain changes when fitted-head and image-similarity uncertainty are examined. The primary contrast was defined before fitting this new run. The study is exploratory because DDI had been examined previously and does not provide released patient identifiers for grouped fitting.')
    m.h('2 Related Work')
    m.p('Fitzpatrick types describe skin response to ultraviolet exposure and should not be interpreted as ethnicity labels {{fst}}. Annotation practice also matters: a study of skin-tone annotation found systematic subjectivity across annotator groups {{schumann}}. We retain the released tone labels rather than infer new labels from images. This choice preserves the benchmark definition but does not remove measurement error. In the current similarity screen, one candidate pair has inconsistent released tone groups, illustrating a source of uncertainty that simple stratification cannot resolve.')
    m.p('Dataset shift in dermatology involves changes in condition prevalence, disease spectrum, acquisition and patient selection. Work on differences in condition distributions across clinical settings shows why adjustment and adaptation can change aggregate model performance {{rikhye}}. The DDI investigators already demonstrated benefits from fine-tuning with diverse DDI images {{ddi}}. Our contribution is a narrower controlled comparison within a frozen representation: matched outcome counts distinguish target-associated additions from general same-source additions, repeated lighter-panel draws expose selection variability, and separate conditional and refitting analyses show which uncertainty each interval includes.')
    m.p('Sigmoid image-text pretraining provides a basis for transferable representations {{siglip}}. The MedGemma technical report and MedSigLIP model card describe the medical image encoder used here {{medgemma,modelcard}}. We do not compare architectures or claim a novel encoder. Freezing the representation allows all classifier arms to use the same image features, but its prior exposure to PAD-UFES-20 remains relevant to the interpretation of development accuracy. The experiment evaluates a specified data-addition procedure rather than a general claim that a particular foundation model removes bias.')
    m.h('3 Materials and Methods')
    m.h('3.1 Study Design and Analysis Status',2)
    m.p('This retrospective secondary study predicts dataset-defined malignant versus non-malignant labels from ordinary clinical photographs. The design does not evaluate referral urgency or prospective patient outcomes. No new diagnostic annotation was created. A machine-readable main analysis plan was stored and SHA256-locked before fitting or inspecting the new classifier scores. It specifies the input hashes, C grid, matching, weighting, resampling and primary V-VI weighted-addition AUROC contrast. This is an execution record, not a registered preregistration or a first-use confirmatory DDI evaluation.')
    m.p('The main plan covers PAD development, baseline DDI transfer, two target-addition arms, matched lighter controls and uncertainty calculations. A separate post hoc plan was declared after the main results were inspected but before similarity-group classifiers were fitted. Its outputs do not overwrite the original folds or scores. Figure 1 shows the training and evaluation boundaries; source versions, hashes and saved split membership are recorded in the supplementary material.')
    m.figure('Figure 1',root/'figures/Figure_M1_study_flow.png','Executed analysis flow. PAD alone determines C. Each adapted head uses all PAD rows and only DDI rows outside its saved evaluation fold. Baseline transfer uses a PAD-only head. The encoder is fixed throughout; resampling refits classifier heads where specified.',width=6.6)
    m.h('3.2 Datasets, Outcomes and Exclusions',2)
    m.p('PAD-UFES-20 version 1 contains 2,298 source images collected with smartphone devices {{pad,padrelease}}. We quarantined all 30 images in 15 exact-byte-duplicate groups rather than retain an arbitrary member where diagnosis or metadata were inconsistent. The curated pool contains 2,268 images from 1,370 source patient identifiers. BCC, SCC and MEL are positive; ACK, SEK and NEV are negative. Bowen disease follows its source grouping with SCC. This is a malignancy mapping: classifying actinic keratosis as negative does not imply that it needs no clinical management. Source diagnoses and patient identifiers were retained.')
    m.p('All 656 DDI images were retained. The released malignant Boolean value defines the outcome. The source DDI study uses pathology-confirmed diagnoses {{ddi}}; patient identifiers are not present in the working release used for this analysis. DDI_ID and file names identify images and are never treated as patient identifiers. Tone labels are grouped as I-II, III-IV and V-VI. No outcome, tone label or score was imputed. PAD has 804 missing tone labels, all attached to negative released outcomes; these rows remain in training and pooled development but do not enter tone-specific descriptions. Table 1 reports the curated counts. Disease counts and exclusions are in Table S2.')
    cap=m.table('Table 1','Curated image cohorts and released tone groups.',['Cohort or tone group','Images','Positive','Negative'],
        [[('PAD-UFES-20' if r['dataset']=='PAD-UFES-20' else 'DDI')+' '+('all' if r['skin_tone_group']=='ALL' else r['skin_tone_group']),r['image_count'],r['positive_count'],r['negative_count']] for r in counts],
        [3.35,1.1,1.1,1.1],'inputs/preparation_counts/dataset_counts.csv')
    m.p('PAD tone groups can contain different images from the same source patient, so per-tone patient counts are not additive. No exact-byte duplicate was found between retained PAD and DDI images. Exact hashing does not exclude photographs of the same person or less-similar repeated views. All classifications and performance estimates are image-level.')
    m.h('3.3 Frozen Features and Classifier Development',2)
    m.p('We used google/medsiglip-448 at revision 9cea28a1a1195f665105faa6e8544c112fd960a4. PIL decoded each image to RGB, discarding any alpha channel, without EXIF transposition or augmentation. The recorded PIL processor resized images to 448 by 448 using its bicubic configuration, rescaled pixel values and applied mean and standard deviation 0.5 per channel. We extracted the image-tower pooler output and L2-normalised it to 1,152 float32 features. Features were cast to float64 for fitting, with no further scaling. The encoder was frozen and no DDI label was used to modify its parameters.')
    m.p('The classifier was L2 logistic regression with an intercept, lbfgs, l1_ratio=0 and tolerance 1e-7, implemented in scikit-learn 1.8.0 {{sklearn,sklearndoc}}. An initial 4,000-iteration budget was increased only through same-objective retries if convergence warnings occurred; unresolved convergence would stop the run. Explicit outcome-class weights were the total unweighted training-row count divided by twice the row count in that class. The software class_weight argument was None. Multiplying these explicit factors by source factors avoids a version-dependent implicit balancing rule. In bootstrap fits the row counts include resampling multiplicities.')
    m.p('Saved PAD folds were created with StratifiedGroupKFold using source patient identifiers. For nested development, each outer training set was sorted by image identifier and split into five grouped inner folds. C was selected from 0.001, 0.01, 0.1, 1, 10, 100 and 1000 by the highest mean inner-fold AUROC, with ties within 1e-12 resolved toward the smallest C. Outer held-out scores provide the nested development estimate. Separately, five-fold grouped development on all PAD selected the C used for all DDI arms; the selected value was 10. Neither DDI outcomes nor DDI scores were used for C selection.')
    m.comment(cap,'The old 1,164-PAD plus 1,059-Fitzpatrick17k pool is replaced by the actual curated 2,268-PAD pool. All downstream tables and claims have been revised to this training definition; these values are not a verification of the old experiment.')
    m.h('3.4 Target Addition and Matched Controls',2)
    m.p('The baseline head was fitted on all PAD images and applied to all DDI images. Adaptation used the saved five DDI image folds, stratified jointly by tone and binary outcome. For fold k, all PAD rows were combined with V-VI DDI rows outside fold k. Every DDI row in fold k was excluded from that fit, including rows of other tones. The unweighted target arm assigned added rows a source factor of one. The weighted arm assigned them the square root of the PAD row count divided by the added-row count. PAD rows retained source factor one. Each effective sample weight was the product of source factor, explicit class factor and any bootstrap multiplicity.')
    m.p('The lighter control drew from I-II and III-IV DDI rows outside the same evaluation fold. Each panel matched the target addition separately on positive and negative counts and sampled without replacement. Twenty draws were generated by independent recorded seeds for each fold; weighted and unweighted arms used exactly the same selected rows. Thus there are 100 lighter training panels, each used for both weighting conditions. The target additions contain 165 or 166 rows per fold, with mean 165.6. This is the quantity evaluated, not a minimum data requirement. Table S6 records fold counts and source multipliers.')
    m.p('Matched target-versus-lighter comparisons have equal added-row counts, outcome counts and source multipliers. Weighted-versus-unweighted comparisons additionally change the total weighted loss and hence the effective influence of the L2 penalty; they are comparisons of specified training procedures. Outcome-count matching does not match diagnoses, body sites, age, lesion severity or acquisition. The target-minus-lighter contrast therefore measures a target-associated increment under this design and does not isolate a causal pigmentation effect.')
    m.h('3.5 Metrics and Uncertainty',2)
    m.p('The primary metric is V-VI AUROC change for weighted target addition relative to PAD-only transfer. Secondary summaries include arm AUROCs, matched contrasts and sensitivity at the most permissive attainable retrospective threshold with empirical FPR at or below 0.50. Negative scores at the boundary are excluded by a strict score-greater-than-threshold rule, preserving ties. Thresholds are estimated from labelled evaluation scores and are not deployment thresholds. We pool held-out logits across DDI folds to compute each AUROC. Lighter-arm summaries average the 20 per-draw metrics; they are not the AUROC of averaged scores.')
    m.p('Conditional uncertainty uses 10,000 shared DDI bootstrap count vectors stratified by original fold, tone and outcome. The same resampled images are used across arms and lighter draws for paired differences. PAD development uses 10,000 patient-cluster bootstrap draws. Percentile 95% intervals follow the bootstrap framework {{efron}}. Conditional intervals hold the fitted heads, selected panels, C, encoder and folds fixed. The variability across 20 lighter panels is reported separately as mean, sample standard deviation and range; it is not 20 independent evaluation cohorts.')
    m.p('A second analysis uses 500 paired refitting bootstrap replicates. PAD patient groups are sampled as clusters, and DDI images are sampled within the same fold, tone and outcome strata. For each replicate we fit one PAD-only head and five weighted-target heads, excluding each test fold from adaptation. DDI multiplicities are shared between the appropriate training and evaluation parts of the replicate. C, the encoder and fold assignments stay fixed. The original point estimate and the mean over refitted replicates are reported separately. These percentile intervals include the specified resampling and head-refitting steps, but omit C-selection, fold-selection, encoder and unobserved DDI patient uncertainty. No multiple-comparison adjustment or equivalence margin is used.')
    m.h('3.6 Post Hoc Similarity Grouping and Traceability',2)
    m.p('We screened verified raw DDI images with a 64-bit DCT perceptual hash: RGB images were converted to grayscale, resized to 32 by 32 with Lanczos, transformed with an orthonormal type-II DCT, and thresholded in the upper 8 by 8 block using the median excluding its DC term. Hamming distance at most four flagged candidates. This screen is diagnostic rather than patient identification. Candidate connected components were grouped, assigned to their minimum original fold and used to refit the two target-addition conditions. Singleton fold assignments remained unchanged. We reused the original PAD-only scores and drew 10,000 uniform component bootstrap samples shared across arms. These conditional component intervals belong to a separate post hoc sensitivity analysis.')
    m.p('The artefact archive links each image identifier and source hash to feature rows, original folds, added training rows, fitted model IDs and logits. It retains coefficient and intercept arrays, C, class order, effective weights and training multiplicities for all heads, including refits. Numerical verification independently reconstructs held-out logits, checks training/evaluation separation and compares bootstrap AUROC and ROC metrics with literal resampling. Figure code reads saved result tables and exports plotted-data CSV files, PNG/SVG graphics and source/output hashes. Reporting details are organised around applicable transparency principles of CLAIM 2024 {{claim}}; patient independence and clinical calibration remain unresolved.')
    m.h('4 Results')
    m.h('4.1 PAD Development and DDI Transfer',2)
    pad_result=data['Table_R2_PAD_development_DDI_transfer.csv'][0]
    m.p(f'Nested patient-grouped PAD development AUROC was {n(pad_result["AUROC"])} (95% CI {b(pad_result["AUROC_CI_low"])} to {b(pad_result["AUROC_CI_high"])}). PAD is named in the encoder pretraining resources, so this estimate is not evidence of independence from encoder training. The final development choice was C=10; the five nested outer-fold choices were 10, 100, 10, 10 and 10. PAD-only DDI AUROC was 0.651 overall, 0.638 at I-II, 0.746 at III-IV and 0.508 at V-VI. Table 2 and Figure 2 show the estimates and their conditional intervals. No PAD V-VI performance estimate is emphasised because that subgroup has only 11 images.')
    m.table('Table 2','Nested PAD development and PAD-only DDI transfer with conditional 95% intervals.',['Evaluation population','Images','Positive','AUROC\n95% CI'],
        [[('PAD nested development' if r['population']=='PAD_nested_development' else 'DDI '+('all' if r['stratum']=='ALL_DDI' else r['stratum'])),r['n'],r['positive'],interval(r['AUROC'],r['AUROC_CI_low'],r['AUROC_CI_high'])] for r in data['Table_R2_PAD_development_DDI_transfer.csv']],
        [3.2,.7,.8,1.95],'results/Table_R2_PAD_development_DDI_transfer.csv')
    m.p('PAD intervals resample source patients and DDI intervals resample images within saved fold, tone and outcome strata. Differences between development and DDI combine changes in image source, disease spectrum, labels and acquisition. The V-VI point estimate is the lowest DDI subgroup value, but a pointwise ordering is not a causal test of skin tone.')
    m.figure('Figure 2',root/'figures/Figure_R1_transfer.png','Nested PAD development and baseline DDI AUROC. Error bars are conditional 95% bootstrap intervals using 10,000 draws. PAD pretraining overlap limits development interpretation. DDI intervals use image resampling and do not establish patient independence.',width=6.4)
    m.h('4.2 Controlled Addition at V-VI',2)
    m.p(f'All arm summaries evaluate the same 207 V-VI images, with 48 positive and 159 negative labels. Unweighted V-VI addition increased AUROC to {n(arm("V-VI","target_unweighted")["AUROC"])}. Weighted addition increased it to {n(primary["left_value"])}; the paired gain was {n(primary["difference_unrounded"],True)} ({b(primary["CI_low"])} to {b(primary["CI_high"])}). Matched lighter additions yielded mean AUROC 0.526 without source weighting and 0.567 with weighting. Arm intervals and operating-point sensitivity are in Table 3; paired comparisons are in Table 4 and Figure 3.')
    rows=[]
    for name in ARM_ORDER:
        a=arm('V-VI',name)
        rows.append([ARM_LABELS[name]+('\nMean of 20 draws' if name.startswith('lighter') else ''),interval(a['AUROC'],a['AUROC_CI_low'],a['AUROC_CI_high']),n(100*float(a['sensitivity_FPR_le_0_50']),digits=1)])
    m.table('Table 3','V-VI arm metrics conditional on fitted heads.',['Training arm','AUROC\n95% CI','Sensitivity at\nFPR <= 0.50 (%)'],rows,[3.25,1.7,1.7],'results/Table_R3_DDI_arms.csv; stratum V-VI')
    m.p('The original V-VI operating point has 79 false-positive images out of 159 negatives, or FPR 0.497. Baseline sensitivity is 22/48, whereas weighted-target sensitivity is 42/48. This threshold uses evaluation labels; these sensitivities cannot be transferred directly to a clinical referral service. Their conditional intervals and paired differences are in Table S10.')
    rows=[]
    for left,right in contrast_pairs:
        a=contrast('V-VI',left,right)
        rows.append([ARM_LABELS[left]+' minus '+ARM_LABELS[right],n(a['difference_unrounded'],True),b(a['CI_low'])+' to '+b(a['CI_high'])])
    cap=m.table('Table 4','Paired controlled AUROC contrasts at V-VI with 10,000 conditional bootstrap resamples.',['Comparison','AUROC change','Paired 95% CI'],rows,[3.95,1.05,1.65],'results/Table_R4_matched_controls.csv; stratum V-VI, metric AUROC')
    m.comment(cap,'All seven rows now contain exact source-derived estimates and intervals, replacing qualitative and approximate historical controls. Lighter additions match positive and negative counts and use the same 20 panels under both weighting conditions.')
    m.p(f'Weighted V-VI addition exceeded weighted lighter addition by {n(target_lighter["difference_unrounded"],True)} ({b(target_lighter["CI_low"])} to {b(target_lighter["CI_high"])}). The corresponding unweighted contrast was +0.159 (0.102 to 0.219). These contrasts support a target-associated increment in this experiment. Their matching does not eliminate differences in within-class diagnoses or patient composition.')
    m.figure('Figure 3',root/'figures/Figure_R2_controls.png','V-VI paired AUROC contrasts and conditional 95% intervals. The first four rows compare the named addition with PAD only. Remaining rows compare two additions. Lighter-arm contrasts average the per-draw metrics on common held-out images.',width=6.4)
    m.h('4.3 Draw Variability and Refitting Robustness',2)
    du=next(r for r in data['matched_draw_summary.csv'] if r['stratum']=='V-VI' and r['arm']=='lighter_unweighted')
    dw=next(r for r in data['matched_draw_summary.csv'] if r['stratum']=='V-VI' and r['arm']=='lighter_weighted')
    m.p(f'The 20 lighter-panel AUROCs ranged from {n(du["AUROC_minimum"])} to {n(du["AUROC_maximum"])} without source weighting, with sample SD {n(du["AUROC_SD"],digits=4)}. With source weighting they ranged from {n(dw["AUROC_minimum"])} to {n(dw["AUROC_maximum"])} with SD {n(dw["AUROC_SD"],digits=4)}. The same DDI images are evaluated in every draw. These ranges describe training-panel variability; they do not represent repeated external validation. Table S7 and Figure S1 provide the individual draws.')
    rows=[]
    for name,label in [('PAD_baseline','PAD only'),('target_weighted','Weighted V-VI addition'),('target_weighted_minus_PAD_baseline','Paired AUROC change')]:
        a=refit('V-VI',name)
        rows.append([label,n(a['original_point_estimate'],name.endswith('minus_PAD_baseline')),n(a['bootstrap_mean'],name.endswith('minus_PAD_baseline')),b(a['CI_low'])+' to '+b(a['CI_high'])])
    cap=m.table('Table 5','Robustness at V-VI from 500 paired bootstrap replicates refitting classifier heads within existing data.',['Estimate','Original point','Refitted mean','Refitting 95% CI'],rows,[2.65,1.0,1.1,1.9],'results/Table_R5_refitting_robustness.csv; stratum V-VI, metric AUROC')
    m.comment(cap,'Table 5 no longer presents an unverifiable Fitzpatrick17k external-replication claim. It now reports the actual 500-replicate within-data head-refitting analysis. This change also removes the historical non-reproduction claim from the abstract, results and discussion.')
    m.p(f'The refitting interval for the original paired AUROC gain was {b(uncertainty["CI_low"])} to {b(uncertainty["CI_high"])}. The bootstrap mean gain was {n(uncertainty["bootstrap_mean"],True)}, compared with original point {n(uncertainty["original_point_estimate"],True)}. The mean is a distinct summary of refitted replicates, not a replacement point estimate. C and folds remain fixed, and the finite 500-replicate percentile approximation should be interpreted accordingly. The weighted-target AUROC interval extends down to 0.660 after refitting, despite the conditional arm interval starting at 0.708.')
    m.figure('Figure 4',root/'figures/Figure_R3_uncertainty.png','Uncertainty around the original weighted V-VI AUROC change. Fixed-head conditional resampling, head-refitting resampling and the post hoc similarity-component analysis use different uncertainty definitions. All remain within existing PAD/DDI data.',width=6.4)
    m.h('4.4 Similarity Sensitivity and Numerical Verification',2)
    m.p('The raw-image screen flagged four DDI candidate pairs at DCT-hash distance two; none had identical decoded RGB pixels. Three pairs crossed original image folds. Grouping their connected components produced 652 similarity components and moved three images to their component minimum folds. One pair has I-II and III-IV tone labels despite near-identical photographs; all four pairs have matching released malignant labels. Original labels were retained. These groupings do not identify patients.')
    m.p('The V-VI baseline and weighted-addition point estimates remained 0.508 and 0.783 because no flagged pair involved V-VI. The component-bootstrap interval for their difference was 0.175 to 0.379. Table S13 reports this new post hoc protocol across tones: the intervention is V-VI addition in every row, rather than a different tone-specific intervention for each row. The component interval remains conditional on the fitted heads and cannot resolve unobserved same-patient photographs.')
    m.p(f'Verification reconstructed {verification["saved_heldout_and_tuning_logits_reconstructed"]:,} saved held-out and tuning logits from {verification["classifier_heads_with_coefficients_and_fit_weights_verified"]:,} fitted heads. It confirmed image order, coefficients, effective weights, multiplicities, excluded test folds, PAD patient-group separation and exact matched outcome counts. Bootstrap metrics were checked against literal repeated-image sampling and scikit-learn ROC/AUROC. The verified archive supplies per-image predictions, all head parameters and numerical samples, so the reported differences are calculated before display rounding.')
    m.h('5 Discussion')
    m.p('Under the specified frozen-encoder procedure, V-VI DDI images show weak PAD-only discrimination and benefit more from V-VI additions than from matched lighter-tone additions. The weighted target-minus-lighter interval excludes zero under conditional resampling. The gain also remains positive when source weighting is removed and when heads are refitted in bootstrap samples. These are consistent observations within one reused clinical cohort; they do not establish transfer to another institution or a causal effect of pigmentation.')
    m.p('The lighter controls help interpret the gain. Matching added outcome counts rules out a simple difference in positive-versus-negative row composition as the explanation for the target-minus-lighter comparison. Weighted lighter addition still yields a smaller positive gain, consistent with a general benefit from same-source images. The additional target-associated increment may reflect tone, diagnoses, body sites, acquisition or other correlated attributes. The design measures their combined association and cannot select one mechanism. Disease-conditioned matching and new patient-grouped cohorts are needed to examine these alternatives.')
    m.p('The separation of conditional and refitting intervals is consequential. Fixed-head intervals quantify evaluation-image uncertainty around the executed classifier scores. Refitting intervals additionally include the recorded changes in training sample composition and fitted coefficients, while still fixing model selection and the encoder. The refitted mean gain is lower than the original point, and the adapted-arm interval becomes wider toward lower AUROC. Reporting both summaries avoids treating a favourable fitted configuration as independent validation. Neither calculation supplies the missing uncertainty from unknown DDI patient overlap.')
    m.h('5.1 Limitations and Generalisability',2)
    m.p('DDI patient independence is the main unresolved limitation. Excluding an image fold prevents its listed rows from entering a corresponding fit, but cannot prevent another photograph of the same patient from entering a different fold. Grouping near-identical images addresses only the flagged similarity pattern. It does not address dissimilar views, lesion-level dependence or missing patient mappings. The adaptation estimate and its intervals may therefore be optimistic. Independent patients with identifiers for grouping are necessary for an external test.')
    m.p('Encoder provenance also limits independence. PAD-UFES-20 is named in MedSigLIP pretraining resources {{modelcard}}; patient grouping during head development does not undo that exposure. Full image-level pretraining overlap with DDI cannot be certified from the public model information. Freezing the encoder controls it between arms, but this does not prove clean unseen-image evaluation. Only one encoder and one C-selection procedure are evaluated, so architecture independence and cross-encoder reproducibility are not established.')
    m.p('Missing tone labels and source outcome definitions constrain interpretation. All 804 missing PAD tone labels belong to negative outcomes, so the distribution of labelled PAD tone groups is not a complete representation of the full training pool. Labels are retained rather than harmonised by inferred tone. Negative-class actinic keratoses and clinically diagnosed benign PAD lesions also distinguish the task from biopsy-confirmed malignancy classification in DDI. We do not equate the binary outcome with referral need or source prevalence with deployment prevalence.')
    m.p('The analysis pools uncalibrated cross-fitted logits from multiple heads; fold-specific intercepts and weighting can affect cross-fold rankings. The pooled AUROC is the reported cross-fitting summary, not a guarantee for one deployable adapted classifier. No probability calibration, prospectively selected threshold, decision-curve analysis, clinician comparison or prospective outcome evaluation was performed. Sensitivity at FPR at most 0.50 is a retrospective analytical summary and represents a high false-positive burden. Multiple secondary comparisons have unadjusted intervals, and neither a familywise confirmatory claim nor performance equivalence is established.')
    m.h('5.2 Implications for Further Evaluation',2)
    m.p('The next external study should test the locked PAD-only and target-adapted procedures on a clinical cohort assembled independently of DDI, with patient identifiers, consistent reference standards and recorded lesion diagnoses and body sites. A design varying the number of target patients can estimate a learning curve. Matching condition distributions as well as outcome counts would help determine how much of the target-associated gain persists after disease-spectrum differences are reduced. Comparing encoders is a separate experiment rather than a conclusion from this run.')
    m.p('For clinical research, threshold selection should use a separate development population and a prospectively defined service constraint. Calibration and patient-level outcomes then need evaluation on independent cases. The current reproducibility archive makes the numerical comparison reviewable without new GPU extraction: each result can be traced to feature rows, folds and a stored classifier head. This traceability supports scrutiny of the experiment, while external effectiveness remains an open empirical question.')
    m.h('6 Conclusions')
    m.p(f'In this exploratory PAD-to-DDI analysis, weighted V-VI addition increased V-VI AUROC from {n(primary["right_value"])} to {n(primary["left_value"])}. The paired change was {n(primary["difference_unrounded"],True)}, with conditional 95% interval {b(primary["CI_low"])} to {b(primary["CI_high"])} and head-refitting interval {b(uncertainty["CI_low"])} to {b(uncertainty["CI_high"])}. Weighted target addition exceeded the outcome-count-matched lighter addition by {n(target_lighter["difference_unrounded"],True)} ({b(target_lighter["CI_low"])} to {b(target_lighter["CI_high"])}). The new similarity-group sensitivity retained the direction of this gain. These results motivate independent patient-grouped testing of target-domain additions; they do not establish a causal skin-tone mechanism, a required minimum sample size, equal performance across tones or clinical effectiveness.')
    m.p('Supplementary Materials: The accompanying document contains Appendices S1-S4, Tables S1-S13 and Figure S1. It provides source-preserving cohort and fold counts, PAD tuning, matched panels, full DDI arm results, refitting summaries, duplicate diagnostics, the separate grouped sensitivity and the artefact/software inventory.')
    for key in ('Author Contributions','Funding','Institutional Review Board Statement','Informed Consent Statement'):
        text=profile['preserved_backmatter'][key]
        if key=='Author Contributions':text=text.replace(' All authors have read and agreed to the published version of the manuscript.','')
        m.p(text,size=9.5)
    m.p('Data Availability Statement: PAD-UFES-20 version 1 is available at DOI 10.17632/zr7vgbcyr2.1, and DDI access is provided by its official dataset portal (https://ddi-dataset.github.io/) under the custodian\'s research-use conditions. The corresponding author retains the versioned numerical archive containing image/feature indices, saved folds, matched-panel membership, pinned model provenance, fixed features, fitted linear-head parameters, per-image logits, locked plans, numerical bootstrap samples and exact analysis/figure scripts. Materials can be provided to editors or researchers subject to applicable source-access and redistribution conditions. Raw DDI photographs and the pretrained encoder checkpoint are not redistributed in this archive. The source encoder is identified by repository revision and recorded file hashes.',size=9.5)
    ack=m.p('Acknowledgments: OpenAI ChatGPT and Codex (accessed August-October 2026) assisted with manuscript editing, Python implementation and execution, numerical verification, scientific plotting and document preparation. Anthropic Claude assisted with earlier language editing and formatting. The authors remain responsible for validating the code, scientific interpretation and final manuscript.',size=9.5)
    m.comment(ack,'The earlier claim that generative AI was never used to assist statistical execution has been corrected. The new run used Codex to implement and execute real numerical analysis; results are backed by stored inputs, scores, models and verification logs.')
    m.p(profile['preserved_backmatter']['Conflicts of Interest'],size=9.5)
    m.h('Abbreviations')
    m.p('AUROC, area under the receiver operating characteristic curve; BCC, basal cell carcinoma; CI, confidence interval; CV, cross-validation; DCT, discrete cosine transform; DDI, Diverse Dermatology Images; FPR, false-positive rate; MEL, melanoma; SCC, squamous cell carcinoma. Table S1 defines the analysis terms.')
    # All supplementary citations reuse the manuscript reference numbering.
    s=Article(profile['references'],order,supplement=True)
    s.p('Supplementary Materials',size=10,bold=True)
    s.p(TITLE,size=17,bold=True,style='Title',keep=True)
    s.p(profile['frontmatter'][0],size=10)
    s.p('This supplement documents the two-source frozen-feature experiment reported in the manuscript. It contains the exact source definitions, fold assignments, tuning summaries, controlled contrasts and uncertainty protocols for the new run. Table S13 uses a separate post hoc similarity-group protocol. Full-precision values, per-image scores and model parameters are retained in the numerical archive; tables here use display rounding.')
    s.h('Appendix S1 Cohorts and Analysis Definitions')
    s.table('Table S1','Definitions used in the article.',['Term','Definition'],[
        ['PAD-only baseline','One head fitted on all curated PAD rows; no DDI fitting rows.'],['Target addition','V-VI DDI rows outside each saved test fold added to all PAD.'],['Lighter control','I-II/III-IV rows matched separately for positive and negative added counts.'],['Conditional interval','Evaluation resampling while fitted heads, C, folds and panels remain fixed.'],['Refitting interval','PAD patient and DDI image resampling plus refitted heads; encoder, C and folds fixed.'],['Similarity component','Connected candidate DCT-hash pairs; a photographic proxy, never a patient identifier.'],['Sensitivity at FPR <= 0.50','Retrospective empirical ROC point, with ties excluded at its score boundary.'],['Mean over 20 draws','Arithmetic mean of per-draw metrics on the same evaluation images.']],
        [1.9,4.75],'plan/MAKE_Analysis_Plan_v1.json; plan/MAKE_Proxy_Group_Plan_v1.json')
    pad_rows=read_csv(root/'inputs/manifests/PAD_labeled_manifest.csv')
    # Diagnosis is preserved in the manifest; field naming follows the actual file.
    diagnosis_field=next(k for k in pad_rows[0] if k in ('source_diagnosis','diagnosis','diagnostic'))
    diag_counts={name:sum(r[diagnosis_field]==name for r in pad_rows) for name in ['BCC','SCC','MEL','ACK','SEK','NEV']}
    s.table('Table S2','PAD diagnoses and source-preserving inclusion rules.',['Source diagnosis or cohort','Rows','Binary definition or exclusion'],
        [[name,diag_counts[name],'Positive' if name in ['BCC','SCC','MEL'] else 'Negative'] for name in diag_counts]+
        [['PAD source release',2298,'Before exact-duplicate quarantine'],['PAD duplicate quarantine',30,'All rows in 15 exact-byte groups excluded'],['PAD retained',2268,'1,370 released source patient identifiers'],['DDI retained',656,'171 malignant and 485 non-malignant; all released rows']],
        [2.1,.65,3.9],'inputs/manifests/PAD_labeled_manifest.csv; source preparation audit')
    s.p('The duplicate quarantine avoids selecting one member of a group with conflicting diagnosis or ambiguous metadata. Missing PAD skin tone was not imputed; 804 such rows are retained and all carry negative labels. PAD cancer diagnoses have a biopsy reference, while benign labels can follow clinical assessment in the source release {{pad,padrelease}}. DDI uses the released pathology-referenced malignant Boolean {{ddi}}. No clinical urgency label was constructed.')
    fold_all=[r for r in folds if r['skin_tone_group']=='ALL']
    s.table('Table S3','Saved outer/evaluation folds before any new classifier fit.',['Dataset','Fold','Images','Positive','Negative','Split units'],
        [[r['dataset'],r['cv_fold'],r['image_count'],r['positive_count'],r['negative_count'],r['split_unit_count']] for r in fold_all],
        [1.75,.6,.85,.9,.9,1.65],'inputs/preparation_counts/fold_counts.csv')
    s.p('Preparation used seed 20261004, source patient StratifiedGroupKFold for PAD and StratifiedKFold on joint tone/outcome classes for DDI. Fold numbering is zero-based. The original saved assignments were not reshuffled after model results were inspected. Full tone-by-fold counts are retained as CSV. DDI split-unit counts refer to images; they do not establish unique patients.')
    s.h('Appendix S2 Development and Controlled Additions')
    tuning=data['PAD_tuning_fold_AUROC.csv']
    mean_rows=[]
    for C in plan['classifier']['C_grid']:
        values=[float(r['AUROC']) for r in tuning if r['scope']=='PAD_final_CV' and float(r['C'])==float(C)]
        assert len(values)==5
        mean_rows.append([n(C,digits=3),n(np.mean(values),digits=6),'Selected' if float(C)==float(fit['C_selected_from_PAD_only']) else ''])
    s.table('Table S4','C selection on full PAD patient-grouped development folds.',['C','Mean fold AUROC','Selection'],mean_rows,[1.05,2.5,3.1],'results/PAD_tuning_fold_AUROC.csv; scope PAD_final_CV')
    s.p('The highest arithmetic mean of five validation-fold AUROCs determines C, with smallest-C tie resolution at tolerance 1e-12. This tuning summary is not the nested outer held-out performance. DDI outcomes and scores do not enter C selection.')
    nested=read_csv(root/'splits/PAD_nested_inner_folds.csv')
    core_registry=read_csv(root/'models/core_models_registry.csv')
    outer_rows=[]
    for fold in range(5):
        r=next(r for r in core_registry if r['model_id']==f'PAD_nested_outer_{fold}')
        f=next(r for r in fold_all if r['dataset']=='PAD-UFES-20' and int(r['cv_fold'])==fold)
        outer_rows.append([fold,2268-int(f['image_count']),f['image_count'],f['split_unit_count'],n(r['C'],digits=0)])
    s.table('Table S5','Nested PAD outer evaluation and internally selected C.',['Outer fold','Training images','Test images','Test patients','Inner-selected C'],outer_rows,[.85,1.45,1.4,1.45,1.5],'models/core_models_registry.csv; splits/PAD_nested_inner_folds.csv')
    s.p('Inner CV uses five StratifiedGroupKFold partitions of each outer training set, sorted by analysis image identifier. Seeds are 20261004 + 100 + outer fold. Every PAD patient is excluded from its corresponding outer training fit. The nested score file links each evaluated image to the exact outer head.')
    rows=[]
    for fold in range(5):
        f=next(r for r in folds if r['dataset']=='DDI' and r['skin_tone_group']=='V-VI' and int(r['cv_fold'])==fold)
        added=207-int(f['image_count']);pos=48-int(f['positive_count']);neg=159-int(f['negative_count'])
        rows.append([fold,added,pos,neg,n(np.sqrt(2268/added),digits=6)])
    s.table('Table S6','Target and matched lighter training counts outside each DDI test fold.',['Test fold','Added rows','Added positive','Added negative','Weighted source factor'],rows,[.85,1.1,1.3,1.3,2.1],'inputs/preparation_counts/fold_counts.csv; splits/DDI_adaptation_samples.csv')
    s.p('Every matched lighter panel has exactly the positive and negative counts in this table. For each fold and draw the eligible lighter identifiers are sorted, sampled without replacement separately by label, and sorted for fitting. Seeds use NumPy SeedSequence with components 20261004, 200, draw ID and fold ID. Weighted and unweighted lighter conditions share panels. Explicit class factors use total row count divided by twice the class row count before source multiplication. Source weighting also changes total loss weight and effective regularisation.')
    rows=[]
    for draw in range(20):
        items=[next(r for r in data['Table_R6_repeated_draws.csv'] if r['stratum']=='V-VI' and r['arm']==name and int(r['draw_id'])==draw) for name in ['lighter_unweighted','lighter_weighted']]
        rows.append([draw]+[v for r in items for v in [n(r['AUROC'],digits=6),n(r['AUROC_change_from_PAD_baseline'],True,digits=6)]])
    s.table('Table S7','V-VI AUROC for every matched lighter-panel draw.',['Draw ID','Unweighted AUROC','Unweighted change','Weighted AUROC','Weighted change'],rows,[.6,1.55,1.45,1.55,1.5],'results/Table_R6_repeated_draws.csv; stratum V-VI')
    s.p('Unweighted mean AUROC is 0.526219 with SD 0.008766; weighted mean is 0.567093 with SD 0.018176. Change is relative to the identical PAD-only V-VI score vector. Draw spread describes selected training panels on a common evaluation cohort. Conditional mean-arm intervals resample the common evaluation images and hold these 20 panels fixed.')
    s.figure('Figure S1',root/'figures/Figure_R4_repeated_draws.png','V-VI AUROC changes from PAD only across 20 matched lighter-panel draws. Horizontal bars mark per-condition means. The dots share the same evaluation images and are not independent validation cohorts.',width=6.2)
    s.h('Appendix S3 Full Numerical Results and Sensitivity Analyses')
    rows=[]
    for tone in TONE_ORDER:
        for name in ARM_ORDER:
            a=arm(tone,name)
            rows.append([tone.replace('ALL_DDI','All DDI'),ARM_LABELS[name],interval(a['AUROC'],a['AUROC_CI_low'],a['AUROC_CI_high'])])
    s.table('Table S8','All DDI arm AUROCs with 10,000 fixed-head conditional intervals.',['Tone group','Training arm','AUROC\n95% CI'],rows,[.9,3.75,2.0],'results/Table_R3_DDI_arms.csv')
    s.p('All DDI contains 656 images; I-II, III-IV and V-VI contain 208, 241 and 207 images, respectively. Target addition always adds V-VI rows, including when evaluation is reported at another tone. Lighter values are means over 20 per-draw metrics. These conditional intervals do not refit heads.')
    rows=[]
    for tone in TONES:
        for left,right in [('target_weighted','PAD_baseline'),('target_unweighted','PAD_baseline'),('target_weighted','lighter_weighted'),('target_unweighted','lighter_unweighted')]:
            a=contrast(tone,left,right)
            rows.append([tone,ARM_LABELS[left]+' minus '+ARM_LABELS[right],n(a['difference_unrounded'],True),b(a['CI_low'])+' to '+b(a['CI_high'])])
    s.table('Table S9','Selected paired controlled AUROC contrasts across DDI tone strata.',['Tone','Comparison','Change','Paired 95% CI'],rows,[.65,3.4,.9,1.7],'results/Table_R4_matched_controls.csv')
    s.p('Every difference is computed from the full-precision point estimates before display rounding. Bootstrap count vectors are shared across compared arms. The complete CSV additionally contains pooled DDI contrasts, weighting contrasts and operating-point sensitivity changes. Secondary contrasts are unadjusted and are not converted into confirmatory tests.')
    rows=[]
    for name in ARM_ORDER:
        a=arm('V-VI',name)
        change='Reference'
        if name!='PAD_baseline':
            z=contrast('V-VI',name,'PAD_baseline','sensitivity_FPR_le_0_50')
            change=interval(z['difference_unrounded'],z['CI_low'],z['CI_high'],True)
        rows.append([ARM_LABELS[name],interval(a['sensitivity_FPR_le_0_50'],a['sensitivity_FPR_le_0_50_CI_low'],a['sensitivity_FPR_le_0_50_CI_high']),change])
    s.table('Table S10','V-VI retrospective sensitivity at empirical FPR at most 0.50.',['Training arm','Sensitivity\n95% CI','Change from PAD\nPaired 95% CI'],rows,[3.25,1.6,1.8],'results/Table_R3_DDI_arms.csv; results/Table_R4_matched_controls.csv; sensitivity metric')
    s.p('Values here are proportions. The original achieved V-VI FPR is 79/159 in all arms and draws. A strict greater-than comparison at the negative-score boundary excludes ties. Thresholds and attained sensitivity are read on labelled evaluation data; no independent threshold calibration is performed. Full per-draw operating points are saved in the archive.')
    rows=[]
    for tone in TONE_ORDER:
        for name,label in [('PAD_baseline','PAD only'),('target_weighted','Weighted target'),('target_weighted_minus_PAD_baseline','Paired change')]:
            a=refit(tone,name)
            sign=name=='target_weighted_minus_PAD_baseline'
            rows.append([tone.replace('ALL_DDI','All DDI'),label,n(a['original_point_estimate'],sign),n(a['bootstrap_mean'],sign),b(a['CI_low'])+' to '+b(a['CI_high'])])
    s.table('Table S11','AUROC uncertainty after 500 paired head-refitting bootstrap replicates.',['Tone','Estimate','Original point','Bootstrap mean','Refitting 95% CI'],rows,[.75,1.75,1.2,1.25,1.7],'results/Table_R5_refitting_robustness.csv; metric AUROC')
    s.p('Six heads are fitted per replicate: one PAD-only baseline and five weighted V-VI adaptation heads. Resampling PAD patient clusters preserves all images of a sampled patient with multiplicity. DDI image counts preserve original fold, tone and outcome strata and are shared between relevant fit and evaluation steps. Source and explicit class factors are updated for multiplicities. C, feature vectors and fold assignments stay fixed. The original estimate and bootstrap mean represent different summaries; this is not independent external replication. The CSV also includes sensitivity refitting summaries.')
    similarity=data['DDI_similarity_candidate_pairs.csv']
    annotation=data['DDI_similarity_metadata_consistency.csv']
    rows=[]
    ddi_index=read_csv(root/'inputs/features/DDI_feature_index.csv');lookup={r['analysis_image_id']:r for r in ddi_index}
    for pair in similarity:
        left,right=lookup[pair['left_image_id']],lookup[pair['right_image_id']]
        rows.append([pair['left_image_id'].replace('DDI:','')+' / '+pair['right_image_id'].replace('DDI:',''),pair['perceptual_hash_distance'],pair['left_fold']+' / '+pair['right_fold'],left['skin_tone_group']+' / '+right['skin_tone_group'],left['target_label']+' / '+right['target_label']])
    s.table('Table S12','Diagnostic DDI similarity pairs and original metadata.',['Image pair','Hash distance','Original folds','Released tones','Positive labels'],rows,[2.15,.85,.95,1.5,1.2],'results/DDI_similarity_candidate_pairs.csv; inputs/features/DDI_feature_index.csv')
    s.p('All flagged pairs have DCT-hash distance two and matching malignant labels. No pair has identical decoded RGB pixels. Pair 000297.png/000546.png crosses tone groups I-II/III-IV; labels were preserved rather than relabelled after scoring. This screen does not establish patient identities. Its source ZIP and individual image hashes were verified against the extraction records.')
    rows=[]
    for tone in TONES:
        a=next(r for r in data['Table_RS13_proxy_group_changes.csv'] if r['stratum']==tone and r['arm']=='target_weighted' and r['metric']=='AUROC')
        u=next(r for r in data['Table_RS13_proxy_group_arms.csv'] if r['stratum']==tone and r['arm']=='target_weighted')
        rows.append([tone,u['n_proxy_components'],n(a['baseline'],digits=6),n(a['adapted'],digits=6),n(a['difference_unrounded'],True,digits=6),b(a['CI_low'])+' to '+b(a['CI_high'])])
    cap=s.table('Table S13','New post hoc image-similarity grouping for weighted V-VI addition evaluated across tones.',['Tone','Components','PAD AUROC','After V-VI addition','Paired change','Conditional 95% CI'],rows,[.55,.9,1.15,1.25,1.15,1.65],'results/Table_RS13_proxy_group_changes.csv; results/Table_RS13_proxy_group_arms.csv')
    s.comment(cap,'This table replaces the unverifiable historical S13, including the inconsistent 0.720 to 0.739 / +0.021 row. It reports a new protocol with 652 photographic components and V-VI-only addition across evaluation tones. Changes are computed before rounding; it is not a recovery of the historical unrounded estimate.')
    s.p('Candidate connected components were assigned to the minimum of their saved original folds; all singleton assignments stayed unchanged. This yielded 652 components and moved three images. A component spans I-II and III-IV, so per-tone component counts sum to 653 rather than 652. No candidate pair involves V-VI, whose point estimates therefore stay unchanged. Conditional intervals use 10,000 uniform component resamples shared across arms. The plan is post hoc and does not establish patient independence. The original main-fold results remain archived separately.')
    s.h('Appendix S4 Software Provenance and Reproducibility')
    s.p('Feature extraction used Python 3.13.15, PyTorch 2.11.0+cu130, Transformers 5.17.0, CUDA 13.0, Pillow 11.3.0 and a Tesla T4. The pinned model was loaded in float32 with SDPA attention and explicit PIL image processing; TF32 was disabled. No augmentation or EXIF transposition was applied. Output vectors were float32, L2-normalised and 1,152-dimensional. The extraction audit records image and source-archive hashes, processor configuration and software versions. Deterministic-algorithm mode was not enabled, so bitwise re-extraction across different GPU environments is not claimed.')
    s.p('CPU head fitting, resampling and verification used Python 3.12.14, NumPy 2.3.5, SciPy 1.17.0, scikit-learn 1.8.0, pandas 2.2.3 and threadpoolctl 3.6.0, with the analysis plan limiting numerical threads to two. Figures used Matplotlib 3.10.8. The source-image similarity screen used Pillow 12.3.0 on CPU. These are execution environments rather than measured throughput benchmarks.')
    s.p('Primary plan SHA256: '+sha(root/'plan/MAKE_Analysis_Plan_v1.json')+'\nFeature-cache identifier (feature_spec.json SHA256): '+sha(root/'inputs/features/feature_spec.json')+'\nPinned encoder revision: 9cea28a1a1195f665105faa6e8544c112fd960a4',size=9.5).alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.p('Raw source encoder weights are identified by their recorded SHA256 in models/encoder_provenance.json and are not included in the numerical archive. Independently checked source configuration files accompany the saved feature matrices. Verification starts from saved vectors rather than a separate image re-extraction.')
    s.p('The archive uses analysis_image_id to link source-preserving feature indices, patient/image fold membership and prediction CSV rows. A model_id resolves to coefficient row, intercept, C, class order, fit multiplicities and effective weights in the corresponding NPZ file and registry CSV. Refitting files contain their held-out image IDs, two score vectors, model IDs, training multiplicities and DDI bootstrap counts together. This permits verification of each metric against its precise fitted head.')
    s.p('There are 426 core classifier heads, 10 post hoc group heads and 3,000 heads in 500 refitting files. Coefficient-based verification reconstructed 767,168 saved held-out and tuning logits, checked excluded evaluation rows and source-patient separation where available, verified all 100 matched lighter panels and recomputed unrounded contrasts and percentile limits. An independent ROC implementation with literal sampled duplicates was used to check bootstrap metric calculation. These checks establish the stated numerical traceability, not clinical validity or DDI patient independence.')
    s.p('The key directories are inputs/features and inputs/manifests for row order and source hashes; inputs/splits and splits for outer, inner and added-panel membership; models for parameters, registries and provenance; predictions for per-image logits; uncertainty for count vectors and metric samples; plan for the two analysis records; results for full-precision tables; and code and figures for executed scripts and plotted-data records. SHA256SUMS covers packaged bytes. The source image custodians retain their access conditions.')
    s.p('Figures 2-4 and Figure S1 are generated by MAKE_Plot_Results_v1.py from their recorded CSV tables. figure_provenance.json links the plotting code, source-data hashes and PNG/SVG outputs. Figure 1 is an original schematic of the executed plan, with separate study_flow_provenance.json. No model result is inferred from a drawn or generated curve, and no clinical photograph is reproduced in the revised manuscript or supplement.')
    s.p('Numerical reproduction starts with requirements-reanalysis.txt and MAKE_Verify_Results_v1.py, using the archive root as --output. The README gives exact refitting, plotting and document commands. These checks need no GPU; independently repeating the image screen requires the authorised DDI ZIP.')
    # Supplementary citations must not introduce an unnumbered extra reference.
    m.h('References')
    for i,key in enumerate(order,1):
        m.p(str(i)+'. '+profile['references'][key],size=9.5)
    for path,article in [(Path(main_path),m),(Path(supp_path),s)]:
        path.parent.mkdir(parents=True,exist_ok=True)
        article.d.save(path)
    table_records=m.tables+s.tables
    source_hashes={path:sha(root/path) for path in sorted({item['source'].split(';')[0] for item in table_records if (root/item['source'].split(';')[0]).is_file()})}
    audit={'completed_utc':datetime.now(timezone.utc).isoformat(),'scope':'new two-source MedSigLIP analysis; original article structure and author declarations retained',
        'code_sha256':sha(__file__),'source_profile_sha256':sha(root/'inputs/paper_source_profile.json'),
        'primary_plan_sha256':sha(root/'plan/MAKE_Analysis_Plan_v1.json'),
        'references_in_order':order,'source_csv_sha256':source_hashes,'tables':table_records,
        'main_comments':m.comments,'supplement_comments':s.comments,
        'documents_sha256':{str(Path(main_path).name):sha(main_path),str(Path(supp_path).name):sha(supp_path)},
        'original_IRB_and_consent_text_preserved':True,'historical_external_replication_claim_retained':False}
    (root/'logs/manuscript_revision_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
    print('REVISED',str(Path(main_path).resolve()),str(Path(supp_path).resolve()))
    print('Tables:',len(m.tables),'main and',len(s.tables),'supplement;',len(order),'references')


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--analysis-root',required=True);p.add_argument('--main-document',required=True);p.add_argument('--supp-document',required=True)
    args=p.parse_args();build(args.analysis_root,args.main_document,args.supp_document)
