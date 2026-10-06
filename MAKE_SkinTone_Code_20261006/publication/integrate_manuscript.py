"""Integrate verified analyses into the supplied MDPI manuscript, preserving its styles.

Inputs are the supplied main/supplement and the two verified numerical archives.
No fitting, synthetic results, external evaluation or public deposition is performed.
"""
from pathlib import Path
from copy import deepcopy
from decimal import Decimal, ROUND_HALF_EVEN
from io import BytesIO
import json, hashlib, re, shutil, argparse
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Mm,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table

parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--output',type=Path,default=Path(__file__).resolve().parents[1])
parser.add_argument('--base',type=Path,default=Path('analysis/make_extension_20261006/source/extracted/MAKE_Reanalysis_v1'))
parser.add_argument('--extension',type=Path,default=Path('analysis/make_submission_20261006/evidence/MAKE_Extension_20261006'))
parser.add_argument('--main-source',type=Path,default=Path('upload/MAKE_SkinCancer_20261006_MDPI.docx'))
parser.add_argument('--supp-source',type=Path,default=Path('upload/MAKE_SkinCancer_20261002_Supplementary_REVIEWED(1).docx'))
args=parser.parse_args()
ROOT=args.output.resolve();BASE=args.base.resolve();EXT=args.extension.resolve()
MAIN=args.main_source.resolve();SUPP=args.supp_source.resolve()
ST={'body':'MDPI_3.1_text','h1':'MDPI_2.1_heading1','h2':'MDPI_2.2_heading2',
    'cap':'MDPI_4.1_table_caption','note':'MDPI_4.3_table_footer','fig':'MDPI_5.2_figure',
    'figcap':'MDPI_5.1_figure_caption','ref':'MDPI_8.1_references','cell':'MDPI_4.2_table_body'}
LABELS={'PAD_baseline':'PAD-only transfer','target_weighted':'V–VI addition, weighted',
    'target_unweighted':'V–VI addition, unweighted','lighter_weighted':'Outcome-matched lighter, weighted',
    'all_DDI_unweighted':'All-DDI addition, unweighted','all_DDI_source_weighted':'All-DDI addition, weighted',
    'all_DDI_group_reweighted':'All-DDI addition, tone reweighted','all_DDI_budget_matched':'All-DDI, matched label budget',
    'DDI_only_class_balanced':'DDI-only, class balanced','DDI_only_group_balanced':'DDI-only, six groups balanced',
    'DFR_style':'DFR-style last-layer refit','target_budget_25':'25 V–VI images',
    'target_budget_50':'50 V–VI images','target_budget_100':'100 V–VI images',
    'diagnosis_matched_target':'Diagnosis-matched V–VI','diagnosis_matched_lighter':'Diagnosis-matched lighter'}
metrics=pd.read_csv(EXT/'results/all_arm_metrics.csv')
contrasts=pd.read_csv(EXT/'results/all_paired_contrasts.csv')
v=metrics[metrics.cohort=='V-VI'].set_index('arm')
cv=contrasts[(contrasts.cohort=='V-VI')&(contrasts.metric=='AUROC')]
coverage=pd.read_csv(EXT/'results/Table_E3_matching_coverage.csv')
fold_summary=pd.read_csv(ROOT/'tables/within_fold_AUROC_summary.csv')

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def metric(a):
    r=v.loc[a];return f'{r.AUROC:.3f} ({r.AUROC_CI_low:.3f} to {r.AUROC_CI_high:.3f})'
def delta(a,b):
    r=cv[(cv.left_arm==a)&(cv.right_arm==b)].iloc[0]
    return f'{r.difference_unrounded:+.3f} ({r.CI_low:+.3f} to {r.CI_high:+.3f})'
def set_text(p,t):
    p.clear();p.add_run(t)
def insert_before(anchor,t,kind='body'):
    p=anchor.insert_paragraph_before(t,style=ST[kind])
    if kind in ['h1','h2','cap']:p.paragraph_format.keep_with_next=True
    return p
def table_before(doc,anchor,caption,rows,widths):
    if caption:insert_before(anchor,caption,'cap')
    t=doc.add_table(rows=1,cols=len(rows[0]));t.autofit=False
    for j,head in enumerate(rows[0]):t.cell(0,j).text=str(head)
    for row in rows[1:]:
        cells=t.add_row().cells
        for j,val in enumerate(row):cells[j].text=str(val)
    for j,w in enumerate(widths):t.columns[j].width=Mm(w)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(row.cells):
            cell.width=Mm(widths[j])
            for p in cell.paragraphs:
                p.style=ST['cell'];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.keep_with_next=False
                for run in p.runs:run.font.size=Pt(8);run.bold=i==0
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','bottom']:
                border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'B8B8B8');borders.append(border)
            cell._tc.get_or_add_tcPr().append(borders)
    # Align new tables with the MDPI body column rather than the publisher margin.
    indent=doc.styles[ST['body']].paragraph_format.left_indent
    if indent:
        ind=OxmlElement('w:tblInd');ind.set(qn('w:w'),str(indent.twips));ind.set(qn('w:type'),'dxa');t._tbl.tblPr.append(ind)
    anchor._p.addprevious(t._tbl)
    return t
def figure_before(doc,anchor,path,caption,width=130):
    p=insert_before(anchor,'','fig');p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Mm(width))
    insert_before(anchor,caption,'figcap')
def append_node(doc,t,kind='body'):
    p=doc.add_paragraph(t,style=ST[kind]);
    if kind in ['h1','h2','cap']:p.paragraph_format.keep_with_next=True
    return p

# Newly added references have temporary identifiers 16–32, remapped by first citation.
NEWREFS={
16:'Kirichenko, P.; Izmailov, P.; Wilson, A.G. Last Layer Re-Training is Sufficient for Robustness to Spurious Correlations. In Proceedings of the International Conference on Learning Representations (ICLR), 2023. https://doi.org/10.48550/arXiv.2204.02937.',
17:'Le, P.Q.; Schlötterer, J.; Seifert, C. Is Last Layer Re-Training Truly Sufficient for Robustness to Spurious Correlations? arXiv 2023, arXiv:2308.00473. https://doi.org/10.48550/arXiv.2308.00473.',
18:'LaBonte, T.; Muthukumar, V.; Kumar, A. Towards Last-layer Retraining for Group Robustness with Fewer Annotations. In Advances in Neural Information Processing Systems 36, 2023. https://doi.org/10.52202/075280-0509.',
19:'Koh, P.W.; Sagawa, S.; Marklund, H.; Xie, S.M.; Zhang, M.; Balsubramani, A.; Hu, W.; Yasunaga, M.; Phillips, R.L.; Gao, I.; et al. WILDS: A Benchmark of in-the-Wild Distribution Shifts. In Proceedings of the 38th International Conference on Machine Learning; PMLR, 2021; Volume 139, pp. 5637–5664. Available online: https://proceedings.mlr.press/v139/koh21a.html (accessed on 6 October 2026).',
20:'Gulrajani, I.; Lopez-Paz, D. In Search of Lost Domain Generalization. In Proceedings of the International Conference on Learning Representations (ICLR), 2021. https://doi.org/10.48550/arXiv.2007.01434.',
21:'Sagawa, S.; Koh, P.W.; Hashimoto, T.B.; Liang, P. Distributionally Robust Neural Networks for Group Shifts: On the Importance of Regularization for Worst-Case Generalization. In Proceedings of the International Conference on Learning Representations (ICLR), 2020. https://doi.org/10.48550/arXiv.1911.08731.',
22:'Groh, M.; Harris, C.; Soenksen, L.; Lau, F.; Han, R.; Kim, A.; Koochek, A.; Badri, O. Evaluating Deep Neural Networks Trained on Clinical Images in Dermatology With the Fitzpatrick 17k Dataset. In Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition Workshops, 2021; pp. 1820–1828. Available online: https://openaccess.thecvf.com/content/CVPR2021W/ISIC/html/Groh_Evaluating_Deep_Neural_Networks_Trained_on_Clinical_Images_in_Dermatology_CVPRW_2021_paper.html (accessed on 6 October 2026).',
23:'Varoquaux, G.; Cheplygina, V. Machine learning for medical imaging: Methodological failures and recommendations for the future. npj Digit. Med. 2022, 5, 48. https://doi.org/10.1038/s41746-022-00592-y.',
24:'Kapoor, S.; Narayanan, A. Leakage and the reproducibility crisis in machine-learning-based science. Patterns 2023, 4, 100804. https://doi.org/10.1016/j.patter.2023.100804.',
25:'Roberts, M.; Driggs, D.; Thorpe, M.; Gilbey, J.; Yeung, M.; Ursprung, S.; Aviles-Rivero, A.I.; Etmann, C.; McCague, C.; Beer, L.; et al. Common pitfalls and recommendations for using machine learning to detect and prognosticate for COVID-19 using chest radiographs and CT scans. Nat. Mach. Intell. 2021, 3, 199–217. https://doi.org/10.1038/s42256-021-00307-0.',
26:'Collins, G.S.; Moons, K.G.M.; Dhiman, P.; Riley, R.D.; Beam, A.L.; Van Calster, B.; Ghassemi, M.; Liu, X.; Reitsma, J.B.; van Smeden, M.; et al. TRIPOD+AI statement: Updated guidance for reporting clinical prediction models that use regression or machine learning methods. BMJ 2024, 385, e078378. https://doi.org/10.1136/bmj-2023-078378.',
27:'Moons, K.G.M.; Damen, J.A.A.; Kaul, T.; Hooft, L.; Andaur Navarro, C.L.; Dhiman, P.; Beam, A.L.; Van Calster, B.; Celi, L.A.; Denaxas, S.; et al. PROBAST+AI: An updated quality, risk of bias, and applicability assessment tool for prediction models using regression or artificial intelligence methods. BMJ 2025, 388, e082505. https://doi.org/10.1136/bmj-2024-082505.',
28:'Seyyed-Kalantari, L.; Zhang, H.; McDermott, M.; Chen, I.Y.; Ghassemi, M. Underdiagnosis bias of artificial intelligence algorithms applied to chest radiographs in under-served patient populations. Nat. Med. 2021, 27, 2176–2182. https://doi.org/10.1038/s41591-021-01595-0.',
29:'Chen, R.J.; Wang, J.J.; Williamson, D.F.K.; Chen, T.Y.; Lipkova, J.; Lu, M.Y.; Sahai, S.; Mahmood, F. Algorithmic fairness in artificial intelligence for medicine and healthcare. Nat. Biomed. Eng. 2023, 7, 719–742. https://doi.org/10.1038/s41551-023-01056-8.',
30:'Gawlikowski, J.; Tassi, C.R.N.; Ali, M.; Lee, J.; Humt, M.; Feng, J.; Kruspe, A.; Triebel, R.; Jung, P.; Roscher, R.; et al. A survey of uncertainty in deep neural networks. Artif. Intell. Rev. 2023, 56 (Suppl. 1), 1513–1589. https://doi.org/10.1007/s10462-023-10562-9.',
31:'Sagawa, S.; Koh, P.W.; Lee, T.; Gao, I.; Xie, S.M.; Shen, K.; Kumar, A.; Hu, W.; Yasunaga, M.; Marklund, H.; et al. Extending the WILDS Benchmark for Unsupervised Adaptation. In Proceedings of the International Conference on Learning Representations (ICLR), 2022. https://doi.org/10.48550/arXiv.2112.05090.',
32:'Stanford Center for Artificial Intelligence in Medicine and Imaging. MRA-MIDAS: Multimodal Image Dataset for AI-based Skin Cancer. Available online: https://aimi.stanford.edu/datasets/mra-midas-Multimodal-Image-Dataset-for-AI-based-Skin-Cancer (accessed on 6 October 2026). https://doi.org/10.71718/15nz-jv40.'}

def make_graphs():
    figdir=ROOT/'figures'
    # Recreate the original vector schematic from its archived plotting recipe.
    fig,ax=plt.subplots(figsize=(7.0,3.1));ax.set(xlim=(0,10),ylim=(0,6));ax.axis('off')
    boxes=[
        (2.55,4.8,4.9,.82,'Frozen MedSigLIP image features\nPAD 2268 images and DDI 656 images'),
        (.25,3.4,4.55,.86,'PAD patient grouped development\nNested evaluation and final C selection'),
        (5.2,3.4,4.55,.86,'DDI five saved image folds\nStratified by skin tone and outcome'),
        (.25,1.93,4.55,.92,'PAD only baseline\nFit all PAD and score all DDI'),
        (5.2,1.93,4.55,.92,'Adaptation and matched controls\nAll PAD plus other DDI folds only'),
        (2.55,.30,4.9,.95,'Held out DDI scores and paired contrasts\nConditional bootstrap and head refitting')]
    for x,y,w,h,text in boxes:
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.03,rounding_size=.06',facecolor='#F5F7F8',edgecolor='#687680',linewidth=1))
        ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=10,color='black')
    for start,end in [((3.8,4.8),(2.5,4.28)),((6.2,4.8),(7.5,4.28)),((2.5,3.4),(2.5,2.89)),((7.5,3.4),(7.5,2.89)),((4.8,3.7),(5.2,2.75)),((2.5,1.93),(3.8,1.28)),((7.5,1.93),(6.2,1.28))]:
        ax.annotate('',xy=end,xytext=start,arrowprops={'arrowstyle':'->','color':'#687680','lw':1.2})
    fig.tight_layout(pad=.5)
    for ext in ['png','svg']:fig.savefig(figdir/('Figure1_study_flow.'+ext),dpi=600,bbox_inches='tight')
    plt.close(fig)
    for old,new in [('Figure_R1_transfer','Figure2_transfer'),('Figure_R2_controls','Figure3_controls'),('Figure_R3_uncertainty','Figure4_uncertainty'),('Figure_R4_repeated_draws','FigureS1_repeated_draws')]:
        for ext in ['png','svg']:shutil.copy2(figdir/(old+'.'+ext),figdir/(new+'.'+ext))
    for old,new in [('Figure_E1_baseline_comparison','Figure5_baselines'),('Figure_E3_diagnosis_sensitivity','Figure7_diagnosis')]:
        for ext in ['png','pdf']:shutil.copy2(EXT/'figures'/(old+'.'+ext),figdir/(new+'.'+ext))
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'axes.spines.top':False,'axes.spines.right':False})
    fig,ax=plt.subplots(figsize=(6.5,3.3))
    arms=['PAD_baseline','target_budget_25','target_budget_50','target_budget_100','target_weighted']
    full_budget=float(coverage.target_available.mean())
    assert full_budget==165.6
    x=[0,25,50,100,full_budget];y=np.array([v.loc[a].AUROC for a in arms]);lo=np.array([v.loc[a].AUROC_CI_low for a in arms]);hi=np.array([v.loc[a].AUROC_CI_high for a in arms])
    ax.plot(x,y,'o-',color='#1F5A99',lw=1.7);ax.fill_between(x,lo,hi,color='#1F5A99',alpha=.12)
    ax.errorbar(x,y,yerr=[y-lo,hi-y],fmt='none',ecolor='#1F5A99',capsize=3)
    ax.set_xticks(x,['0','25','50','100','All\n165–166']);ax.set_ylim(.35,.92)
    ax.set_xlabel('V–VI images added in each non-test fold');ax.set_ylabel('V–VI AUROC');ax.grid(alpha=.15)
    ax.set_title('Annotation learning curve',loc='left',fontweight='bold');fig.tight_layout()
    for ext in ['png','pdf']:fig.savefig(figdir/('Figure6_learning_curve.'+ext),dpi=600,bbox_inches='tight',facecolor='white')
    plt.close(fig)
    (ROOT/'audit/learning_curve_axis_check.json').write_text(json.dumps({'full_budget_by_fold':coverage.target_available.tolist(),'mean_full_budget':full_budget,'numeric_AUROC_and_CIs_unchanged':True,'axis_source':'Table_E3_matching_coverage.csv target_available'},indent=2)+'\n')

def update_fields(doc):
    node=doc.settings.element.find(qn('w:updateFields'))
    if node is None:node=OxmlElement('w:updateFields');doc.settings.element.append(node)
    node.set(qn('w:val'),'true')

def build_main():
    doc=Document(MAIN);orig=doc.paragraphs.copy();p={i:q for i,q in enumerate(orig)}
    oldrefs={i+1:deepcopy(q._p) for i,q in enumerate(orig[96:111])}
    set_text(p[6],
        'Targeted annotations may improve skin-lesion classification, but benefits can reflect annotation budget or diagnosis composition. We compared adaptation procedures using frozen MedSigLIP features from 2,268 PAD-UFES-20 images and 656 Diverse Dermatology Images (DDI). Classifier development used patient-grouped PAD folds; DDI adaptation and evaluation used five saved image folds. Weighted Fitzpatrick V–VI addition increased V–VI area under the receiver operating characteristic curve (AUROC) from 0.508 to 0.783; the paired gain was +0.275 (95% interval, 0.131–0.357) after 500 head-refitting bootstrap replicates. Exploratory controls evaluated all-DDI addition, group reweighting, DDI-only heads, last-layer refitting, learning curves and diagnosis matching. Using 5,000 conditional image/panel resamples, targeted addition exceeded an equally sized, outcome-matched all-tone addition by +0.131 (0.075–0.190), while a six-group-balanced DDI-only head using more labels achieved AUROC 0.849. AUROC increased from 0.619 with 25 targeted images to 0.736 with 100 and 0.783 with all 165–166 available images per fold. Exact diagnosis matching yielded a target-minus-lighter difference of +0.046 (−0.002–0.095). These results support annotation efficiency within DDI. Unknown patient overlap, encoder pretraining exposure and absent independent external evaluation preclude diagnosis-independent or general superiority claims.')
    set_text(p[7],'Keywords: skin lesion classification; skin tone; domain adaptation; annotation budget; diagnosis matching; MedSigLIP; bootstrap uncertainty')
    set_text(p[11],p[11].text+' The additional controls examine annotation budget and diagnostic composition, which are distinct from aggregate sample size and binary outcome balance.')
    set_text(p[12],
        'We use smartphone clinical images from PAD-UFES-20 (PAD) to develop the classifier [4,5] and the released DDI cohort for transfer evaluation and within-cohort adaptation. MedSigLIP supplies a fixed medical image representation [6]. We ask whether targeted V–VI additions outperform outcome-matched lighter additions, how that comparison changes against stronger same-source and budget-matched baselines, how discrimination changes with annotation count, and how much advantage remains after exact diagnosis matching. The primary reanalysis contrast was recorded before fitting the archived PAD/DDI classifiers; the additional E1–E3 protocol was recorded before its new fits, after the original results were known. Both analyses are exploratory because DDI had been examined previously and lacks released patient identifiers for grouped fitting.')
    set_text(p[15],p[15].text.replace('Our contribution is a narrower controlled comparison within a frozen representation:','Our initial comparison within a frozen representation used the following controls:'))
    for text in [
        'Balanced last-layer retraining provides a strong comparator for subgroup adaptation. Deep feature reweighting (DFR) refits a linear head on a balanced subset of frozen representations [16]; its applicability to medical data has also been examined [17]. Work on annotation-efficient last-layer retraining shows that label requirements are themselves an experimental variable [18]. We therefore include a DFR-style head and an all-tone control with the same annotation budget as targeted addition. These controls do not reproduce the original DFR benchmarks or establish the superiority of a new architecture.',
        'Distribution-shift benchmarks distinguish changes between source and target environments from performance within an already observed cohort [19]. Domain-generalisation research further demonstrates the importance of model-selection rules when comparing methods [20]. Group distributionally robust optimisation targets worst-group risk through a different objective [21]; our fixed group weights are reported as reweighting, rather than described as group-DRO training. Clinical-image work with Fitzpatrick17k provides complementary evidence on skin-type-associated generalisation [22], but that dataset is not included in the executed training pool.'
    ]:insert_before(p[17],text)
    set_text(p[19],p[19].text.replace('before any classifier was fitted or any classifier score was inspected','before fitting or inspecting the classifier scores of the archived PAD/DDI reanalysis'))
    set_text(p[20],p[20].text+' An additional exploratory protocol recorded on 6 October 2026 covers E1–E3 and prespecified external-evaluation heads. It was recorded after the original DDI results were known and is not a first-use preregistration. Its SHA256 and input hashes are retained in Appendix S5.')
    set_text(p[36],p[36].text+' The additional diagnosis-matched control in Section 2.7 addresses the released diagnosis distribution within its common support, while retaining these other limitations.')
    set_text(p[43],p[43].text.replace('according to the locked plans','according to the recorded protocols')+' The 6 October 2026 extension and document revision used Codex based on GPT-6. Model version identifiers for earlier ChatGPT/Claude assistance were not available in the supplied records.')
    set_text(p[43],p[43].text.replace('according to the locked plans','according to the separately recorded analysis protocols'))
    for kind,text in [
        ('h2','2.7. Additional Baselines, Annotation Budgets and Diagnosis Matching'),
        ('body','E1 retained the existing frozen feature cache, PAD-selected L2 penalty C=10, original DDI folds and test exclusions. Full all-DDI additions combined PAD with every DDI image outside the held-out fold, with either unit source weight, the original square-root source factor, or that source factor together with equal tone contribution within each outcome class. Separate DDI-only heads used inverse-class or six tone-by-outcome group weights. Full-DDI controls used 524–525 non-test DDI labels per fold, whereas targeted addition used 165–166. Twenty all-tone panels matched targeted addition in total label budget and positive/negative counts. Available labels, sampling panels and effective weights are recorded for every fitted head.'),
        ('body','The DFR-style comparator used source-only PAD standardisation, an L1 logistic head on equally sized tone-by-outcome samples and the average coefficients and intercept of 20 balanced refits [16]. Its C grid was 0.01, 0.03, 0.07, 0.1, 0.3, 0.7 and 1. Three inner folds of the non-test DDI pool, with four matched balanced draws per inner fold, selected C by mean worst-six-group accuracy at a zero logit threshold, resolving ties toward the smallest C. The full non-test labeled pool was available for tuning and group construction. No outer test image entered fitting or selection. This frozen-MedSigLIP adaptation of DFR differs from the original feature-training experiments.'),
        ('body','E2 used 20 nested, outcome-stratified V–VI panels at budgets of 25, 50 and 100 images per non-test fold. Positive quotas followed the available V–VI outcome proportion with half-up rounding; class and source weights were recomputed at each budget. The all-data endpoint reused the original weighted-target scores at fold budgets of 165, 166, 165, 166 and 166 images (mean 165.6). E3 paired target and lighter panels on exact released diagnosis-by-outcome cells, retaining the smaller available count in each cell. Twenty sampled paired panels per fold had identical diagnoses, outcome counts and added-row budgets. Only non-test diagnoses determined common support; all held-out evaluation images remained included.'),
        ('body','The extension used 5,000 shared fold-by-tone-by-outcome image bootstrap count vectors. Sampled-panel arms additionally resampled their 20 fitted panel realisations, sharing panel counts within the paired E3 and nested E2 comparisons. Point estimates average per-panel AUROCs rather than computing AUROC from averaged random-panel scores. All effects were computed before rounding. The resulting unadjusted exploratory intervals are conditional on the saved heads and folds. They differ from the original 10,000-resample fixed-panel intervals in Tables 2–4 and the 500-refit intervals in Table 5; they neither refit the new baselines nor constitute independent replication.'),
        ('body','As a separately recorded post hoc ranking diagnostic, we also computed AUROC within each original held-out fold and averaged the fold/panel estimates. This descriptive summary avoids comparisons between logits produced by different fold-specific heads. It has a different estimand from pooled AUROC and is not accompanied by an inferential interval or treated as external validation. Table S17 reports the results.')
    ]:insert_before(p[44],text,kind)
    set_text(p[58],p[58].text+' Exact diagnosis matching and stronger baselines are considered separately in Sections 3.5–3.7.')
    # Preserve the positive lower limit without disguising its sign by coarse rounding.
    doc.tables[4].cell(4,2).text='0.00004 to 0.037'
    for q in doc.tables[4].cell(4,2).paragraphs:q.style=ST['cell']
    set_text(p[70],p[70].text+' The extension archive separately retains 1,261 classifier records and 97,744 image-level held-out/transfer prediction records. Its verifier reconstructed all of those logits and independently reconstructed 91,840 inner-validation predictions, 504 DFR tuning summaries and the selected penalty values. Stored identifiers, panels, weights and hash records connect each result to its fitted parameters; the counts describe stored records rather than independent images.')
    e1=['PAD_baseline','target_weighted','target_unweighted','lighter_weighted','all_DDI_unweighted','all_DDI_source_weighted','all_DDI_group_reweighted','all_DDI_budget_matched','DDI_only_class_balanced','DDI_only_group_balanced','DFR_style']
    insert_before(p[71],'3.5. Stronger Baselines and Matched Annotation Budget','h2')
    insert_before(p[71],f'All E1 arms were evaluated on the same 207 V–VI images, including 48 malignant cases. Weighted targeted addition reached {metric("target_weighted")}; the six-group-balanced DDI-only head reached {metric("DDI_only_group_balanced")} using a larger available annotation budget. Table 6 and Figure 5 therefore distinguish discrimination from annotation efficiency.')
    rows=[['Training procedure','DDI labels/fold','AUROC (95% CI)']]
    for a in e1:
        n='0' if a=='PAD_baseline' else ('165–166' if a in ['target_weighted','target_unweighted','lighter_weighted','all_DDI_budget_matched'] else '524–525')
        rows.append([LABELS[a],n,metric(a)])
    table_before(doc,p[71],'Table 6. E1 baselines and annotation budgets on held-out V–VI images.',rows,[60,23,49])
    insert_before(p[71],'Full-DDI methods have a larger available label budget. DFR uses the full pool for tuning/group construction and balanced subsets for final refits. Intervals use the extension image/panel protocol (5,000 replicates). AUROC, area under the receiver operating characteristic curve; CI, confidence interval; DDI, Diverse Dermatology Images; DFR, deep feature reweighting.','note')
    insert_before(p[71],f'Targeted addition exceeded the equally sized, outcome-matched all-tone addition by {delta("target_weighted","all_DDI_budget_matched")}. Its differences against full source-weighted addition ({delta("target_weighted","all_DDI_source_weighted")}), tone-reweighted addition ({delta("target_weighted","all_DDI_group_reweighted")}) and DFR-style refitting ({delta("target_weighted","DFR_style")}) had intervals spanning zero. These intervals do not establish equivalence. The targeted-minus-DDI-only six-group contrast was {delta("target_weighted","DDI_only_group_balanced")}.')
    figure_before(doc,p[71],ROOT/'figures/Figure5_baselines.png','Figure 5. E1 V–VI AUROC and exploratory 95% intervals. Full-DDI controls have 524–525 labels available per fold; targeted and budget-matched additions have 165–166. All estimates use the same held-out images. See Table 6 for the full training procedures.')
    insert_before(p[71],'3.6. Annotation Learning Curve','h2')
    insert_before(p[71],f'V–VI AUROC increased with the evaluated target-label budgets: {v.loc["target_budget_25"].AUROC:.3f} at 25 images, {v.loc["target_budget_50"].AUROC:.3f} at 50, {v.loc["target_budget_100"].AUROC:.3f} at 100 and {v.loc["target_weighted"].AUROC:.3f} with all available images (Table 7; Figure 6). The gain from 100 to all available images was {delta("target_weighted","target_budget_100")}; these evaluated budgets do not establish a plateau or a minimum sufficient sample size.')
    rows=[['V–VI labels/fold','AUROC (95% CI)','Panel AUROC SD']]
    for a,n in [('PAD_baseline','0 (PAD only)'),('target_budget_25','25'),('target_budget_50','50'),('target_budget_100','100'),('target_weighted','All: 165–166')]:rows.append([n,metric(a),f'{v.loc[a].AUROC_panel_SD:.3f}' if a.startswith('target_budget') else 'Single score vector'])
    table_before(doc,p[71],'Table 7. E2 nested annotation learning curve on V–VI images.',rows,[33,58,41])
    insert_before(p[71],'Twenty nested sampled panels per budget; the all-data endpoint uses the saved original targeted model. SD describes panel variability on a common evaluation cohort. Intervals use 5,000 extension image/panel resamples.','note')
    figure_before(doc,p[71],ROOT/'figures/Figure6_learning_curve.png','Figure 6. V–VI annotation learning curve. Bars and shading are conditional image/panel 95% intervals. The all-data position uses the observed mean fold budget of 165.6 images; neither the band nor the 20 panels represents independent external replication.')
    insert_before(p[71],'3.7. Exact Diagnosis Matching','h2')
    insert_before(p[71],f'Exact diagnosis common support retained 114–125 images per source/fold, or {coverage.target_fraction_retained.min()*100:.1f}–{coverage.target_fraction_retained.max()*100:.1f}% of available V–VI adaptation images. The matched target and lighter arms reached {metric("diagnosis_matched_target")} and {metric("diagnosis_matched_lighter")}, respectively. Their paired difference was {delta("diagnosis_matched_target","diagnosis_matched_lighter")} (Table 8). The interval includes zero and meaningful positive effects, so this experiment does not establish a diagnosis-independent advantage or the absence of such an advantage.')
    rows=[['Diagnosis-matched source','V–VI AUROC / change (95% CI)'],[LABELS['diagnosis_matched_target'],metric('diagnosis_matched_target')],[LABELS['diagnosis_matched_lighter'],metric('diagnosis_matched_lighter')],['Target minus lighter',delta('diagnosis_matched_target','diagnosis_matched_lighter')]]
    table_before(doc,p[71],'Table 8. E3 exact diagnosis-matched adaptation on all held-out V–VI images.',rows,[66,66])
    insert_before(p[71],'Each paired fit has equal diagnosis/outcome counts and an equal label budget of 114–125 images/source/fold. The same full V–VI test cohort is retained. See Tables S15–S16 for contrasts and support coverage.','note')
    figure_before(doc,p[71],ROOT/'figures/Figure7_diagnosis.png','Figure 7. Target-minus-lighter AUROC contrasts under outcome matching and exact diagnosis common support. Both intervals use the common 5,000-resample extension protocol, including panel resampling. Adaptation supports and budgets differ between rows, so the reduction in the contrast is not an identified causal effect of diagnosis or skin tone.')
    insert_before(p[71],'3.8. Within-Fold Ranking Diagnostic','h2')
    insert_before(p[71],'The descriptive mean within-fold V–VI AUROCs were 0.514 for PAD only, 0.792 for weighted targeted addition, 0.657 for the matched-budget all-tone addition and 0.849 for the six-group-balanced DDI-only head (Table S17). Diagnosis-matched target and lighter means were 0.700 and 0.647. These summaries preserve the main numerical ordering while avoiding cross-fold logit comparisons; they do not remove patient overlap, sampling uncertainty or the need for an external cohort.')
    set_text(p[72],
        'The additional controls sharpen the interpretation of targeted labeling. Under a fixed representation, weighted V–VI addition improved discrimination relative to PAD-only transfer and an equally sized, outcome-matched all-tone addition. However, a balanced DDI-only head with a larger label budget performed better, and full all-DDI or DFR-style heads were not clearly inferior. The evidence therefore supports an annotation-efficiency result within this dataset, rather than universal superiority of targeted addition. This distinction follows the need for credible baselines and comparable evaluation conditions in medical-imaging methods research [23–25].')
    set_text(p[73],
        'Exact diagnosis matching changed the apparent target-versus-lighter advantage. The original outcome-matched comparison retained a positive conditional interval, whereas the new diagnosis-matched difference was +0.046 with an interval spanning zero. Matching diagnoses also removed non-overlapping training cases and reduced each adaptation budget, so these contrasts are different estimands. Their difference cannot be attributed entirely to disease mix, and the diagnosis-matched interval cannot be interpreted as proof of no effect. The observed attenuation nevertheless requires restraint in assigning the original improvement to skin tone alone.')
    set_text(p[74],p[74].text+' The extension intervals additionally include resampling of the realised training panels, while holding their fitted heads fixed. Stronger-baseline refitting, alternative fold construction and external replication remain unevaluated. Performance intervals should also be distinguished from uncertainty in an individual patient prediction [30].')
    insert_before(p[75],
        'The learning curve indicates continuing gains over the evaluated V–VI budgets, but does not establish the least number of patients needed for a clinical system. Each budget counts labeled images; multiple views of one patient cannot be quantified from released DDI metadata. The available-label cost for a full-DDI comparator also differs from the training subset used by a final balanced head. These results should be read as controlled allocation of annotations, not as a universal rule that target subgroup labels are always preferable.')
    set_text(p[76],p[76].text+' The added diagnosis and budget controls retain the same image folds and therefore inherit this dependence limitation; they do not remedy it.')
    set_text(p[79],p[79].text+' Reporting and risk-of-bias frameworks distinguish transparent analysis from evidence suitable for clinical use [26,27]. Subgroup discrimination alone also does not establish equitable decisions or clinical outcomes [28,29]. No formal PROBAST+AI assessment or full clinical deployment evaluation is claimed.')
    set_text(p[81],
        'The decisive next experiment is external evaluation of the frozen PAD-only, targeted, all-DDI and balanced-head procedures on an independently assembled cohort. MRA-MIDAS supplies clinical and dermoscopic images with patient/lesion metadata and a prospectively recruited source design [32]. The prepared protocol restricts the proposed external analysis to eligible biopsied clinical-image lesions, uses a reviewed malignant-versus-benign endpoint with intermediate diagnoses excluded, averages clinical-view logits by lesion and uses paired patient-cluster uncertainty. Seven deployment heads are already saved, and the evaluator has passed input-validation tests; no MRA-MIDAS features or clinical performance results were available in the present study. Source-use agreement review, endpoint harmonisation and cross-dataset overlap checks are required before this external evaluation can be completed.')
    set_text(p[82],p[82].text+' Labeled last-layer adaptation should also be distinguished from unsupervised adaptation using unlabeled target data, as evaluated in the extended WILDS benchmark [31]. The present experiments evaluate the former setting.')
    set_text(p[84],
        'In this exploratory PAD-to-DDI study, weighted V–VI addition increased V–VI AUROC from 0.508 to 0.783, with an original head-refitting interval of 0.131 to 0.357 for the gain. At an equal annotation budget, targeted addition exceeded an outcome-matched all-tone control by +0.131 (95% CI 0.075 to 0.190). A six-group-balanced DDI-only head using more labels reached AUROC 0.849, and exact diagnosis matching yielded a smaller target-minus-lighter contrast of +0.046 (−0.002 to 0.095). The supported contribution is a controlled comparison of annotation allocation, reweighting and diagnostic composition within a shared frozen representation. Independent patient-grouped external evaluation is needed before claims of clinical effectiveness, a diagnosis-independent skin-tone benefit or general superiority can be made.')
    set_text(p[85],
        'Supplementary Materials: The accompanying document contains Appendices S1–S6, Tables S1–S18 and Figure S1. It reports source-preserving cohorts, folds, tuning, original matched controls, refitting and similarity sensitivities, all extension arm metrics, paired contrasts, diagnosis-support coverage, within-fold ranking diagnostics and the reproducibility/external-evaluation inventory.')
    set_text(p[90],p[90].text+' The separate versioned E1–E3 evidence archive contains the extension protocol, saved classifier parameters, image-level predictions, sampling panels, confidence-interval samples and graph provenance. Access to research artifacts remains subject to the source custodians’ conditions.')
    set_text(p[91],p[91].text.replace('as described in Section 2.6','as described in Sections 2.6–2.7').replace('the prespecified analyses','the separately recorded original and exploratory additional analyses'))
    for label,meaning in [('DFR','Deep feature reweighting'),('L1','Sum-of-absolute-coefficients penalty'),('L2','Sum-of-squared-coefficients penalty'),('MRA-MIDAS','Melanoma Research Alliance Multimodal Image Dataset for AI-based Skin Cancer')]:
        cells=doc.tables[-1].add_row().cells;cells[0].text=label;cells[1].text=meaning
    # Retain original reference typography; remap and order all references by first citation.
    pat=re.compile(r'\[([0-9, –-]+)\]')
    def expand(s):
        nums=[]
        for part in s.replace(' ','').split(','):
            limits=re.split('[–-]',part)
            nums.extend(range(int(limits[0]),int(limits[1])+1) if len(limits)==2 else [int(part)])
        return nums
    order=[]
    for q in doc.paragraphs:
        if q.style.name==ST['ref']:continue
        for match in pat.finditer(q.text):
            for n in expand(match.group(1)):
                if n not in order:order.append(n)
    assert set(order)==set(range(1,33)),('Uncited references',set(range(1,33))-set(order))
    mapping={old:i+1 for i,old in enumerate(order)}
    def remap(match):return '['+','.join(str(mapping[n]) for n in sorted(expand(match.group(1)),key=lambda n:mapping[n]))+']'
    for q in doc.paragraphs:
        if q.style.name!=ST['ref'] and pat.search(q.text):set_text(q,pat.sub(remap,q.text))
    for q in orig[96:111]:q._p.getparent().remove(q._p)
    for old in order:
        if old in oldrefs:node=deepcopy(oldrefs[old]);p[111]._p.addprevious(node)
        else:
            node=deepcopy(oldrefs[15]);new=Paragraph(node,doc);new.clear();new.add_run(NEWREFS[old]);p[111]._p.addprevious(node)
    # Replacing original scientific rasters preserves the MDPI picture paragraphs.
    names=['Figure1_study_flow','Figure2_transfer','Figure3_controls','Figure4_uncertainty']
    for index,name in zip([21,50,59,65],names):
        width=p[index].runs[0]._r.xpath('.//wp:extent/@cx')
        original_width=int(width[0]) if width else int(Mm(130))
        p[index].clear();p[index].add_run().add_picture(str(ROOT/'figures'/(name+'.png')),width=original_width)
    # Keep the publisher float after front matter for faithful PDF conversion.
    publisher_float=doc.tables[0]._tbl.tblPr.find(qn('w:tblpPr'))
    if publisher_float is not None:publisher_float.set(qn('w:tblpY'),'11200')
    p[8]._p.addnext(doc.tables[0]._tbl)
    for t in doc.tables[1:]:
        if len(t.rows)<=13:
            for i,row in enumerate(t.rows):
                for cell in row.cells:
                    for q in cell.paragraphs:q.paragraph_format.keep_with_next=i<len(t.rows)-1
    update_fields(doc)
    doc.core_properties.subject='Complete integration of verified original and exploratory E1–E3 results; external MRA-MIDAS evaluation not performed.'
    destination=ROOT/'MAKE_SkinCancer_20261006_MDPI.docx';doc.save(destination)
    (ROOT/'audit/reference_number_mapping.json').write_text(json.dumps(mapping,indent=2)+'\n')
    (ROOT/'audit/reference_inventory.json').write_text(json.dumps([{'number':mapping[n],'previous_number':n if n<=15 else None,'added':n>15,'reference':Paragraph(oldrefs[n],doc).text if n<=15 else NEWREFS[n]} for n in order],indent=2,ensure_ascii=False)+'\n')
    return doc,mapping,oldrefs

def build_supplement(main,mapping,oldrefs):
    source=Document(SUPP);doc=Document(MAIN)
    for child in list(doc.element.body):
        if child.tag!=qn('w:sectPr'):doc.element.body.remove(child)
    append_node(doc,'Supplementary Materials','h1')
    doc.add_paragraph(main.paragraphs[1].text,style='MDPI_1.2_title')
    doc.add_paragraph('Sakchai Tangprasert and Suwit Chanthasen',style='MDPI_1.3_authornames')
    append_node(doc,'This supplement reports the verified original PAD/DDI reanalysis and the separately recorded exploratory E1–E3 extension. Tables S1–S13 retain their original numerical definitions; Tables S14–S18 document the extension and descriptive ranking audit. No independent MRA-MIDAS performance result is included. All changes and intervals are calculated before display rounding.')
    tables=list(source.tables);table_index=0
    for child in source.element.body:
        if child.tag==qn('w:p'):
            q=Paragraph(child,source);text=q.text
            if text.startswith(('Supplementary Materials','Controlled Data Addition','Sakchai Tangprasert','This supplement documents')):continue
            if q._p.xpath('.//w:drawing'):
                pic=doc.add_paragraph('',style=ST['fig']);pic.paragraph_format.keep_with_next=True
                pic.add_run().add_picture(str(ROOT/'figures/FigureS1_repeated_draws.png'),width=Mm(130));continue
            text=re.sub(r'\[([0-9,]+)\]',lambda m:'['+','.join(str({2:1,4:2,5:3}.get(int(n),int(n))) for n in m.group(1).split(','))+']',text)
            text=text.replace('V-VI','V–VI').replace('I-II','I–II').replace('III-IV','III–IV').replace('FPR <=','FPR ≤')
            kind='h1' if q.style.name.startswith('Heading') else ('cap' if text.startswith('Table S') else 'figcap' if text.startswith('Figure S') else 'body')
            if text:append_node(doc,text,kind)
        elif child.tag==qn('w:tbl'):
            # Retain cell text while rebuilding imported tables as canonical XML.
            source_table=Table(child,source)
            widths=np.array([c.width for c in source_table.columns],dtype=float)
            widths=widths/widths.sum()*132
            rows=[[c.text for c in r.cells] for r in source_table.rows]
            anchor_old=doc.add_paragraph('',style=ST['body'])
            table_before(doc,anchor_old,None,rows,widths)
            anchor_old._p.getparent().remove(anchor_old._p);table_index+=1
    assert table_index==13
    append_node(doc,'Appendix S5. Exploratory Extension: Baselines, Budgets and Diagnoses','h1')
    append_node(doc,'The extension protocol SHA256 is 58a6164a2b2fb03941887679e6cf21ebb1271d5b5522ea1c4ac04e05456f86b9. It was recorded before the additional fitting, after the original DDI results were known, and is exploratory rather than preregistered. The source numerical ZIP has SHA256 bb22bcafdef2e13a85131d503b962dc8089ba96107ae0fcb96f4364fd31bd3a7. DFR-style fitting adapts the balanced last-layer method [5] to the same frozen MedSigLIP representation [4].')
    append_node(doc,'All additional fits reuse the original outer DDI image folds and C=10 for L2 heads. Tone-by-outcome weights, source factors, inner DFR tuning, sample panels and the 5,000 shared image/panel resamples are specified in Section 2.7 of the article and the machine-readable protocol. A sampled-arm point is the mean of per-panel metrics, not a metric on averaged sampled-arm scores. Extension intervals do not refit their heads; original Tables S8–S11 use their separate uncertainty protocols.')
    anchor=doc.add_paragraph('',style=ST['body'])
    rows=[['Evaluation cohort','Training procedure','AUROC (95% CI)']]
    for cohort in ['ALL_DDI','I-II','III-IV','V-VI']:
        for _,r in metrics[metrics.cohort==cohort].iterrows():rows.append([cohort.replace('-','–'),LABELS[r.arm],f'{r.AUROC:.3f} ({r.AUROC_CI_low:.3f} to {r.AUROC_CI_high:.3f})'])
    table_before(doc,anchor,'Table S14. All 64 extension arm/cohort AUROC estimates with exploratory 95% intervals.',rows,[22,58,52])
    insert_before(anchor,'ALL_DDI includes 656 images; I–II, III–IV and V–VI include 208, 241 and 207, respectively. Sampled arms use 20 panel realisations. The numerical CSV retains every unrounded value and secondary operating-point metric.','note')
    pairs=[('target_weighted','PAD_baseline'),('target_weighted','lighter_weighted'),('target_weighted','all_DDI_unweighted'),('target_weighted','all_DDI_source_weighted'),('target_weighted','all_DDI_group_reweighted'),('target_weighted','all_DDI_budget_matched'),('target_weighted','DDI_only_class_balanced'),('target_weighted','DDI_only_group_balanced'),('target_weighted','DFR_style'),('target_budget_100','target_budget_50'),('target_weighted','target_budget_100'),('diagnosis_matched_target','diagnosis_matched_lighter')]
    rows=[['V–VI contrast','AUROC difference (95% CI)']]+[[LABELS[a]+' minus '+LABELS[b],delta(a,b)] for a,b in pairs]
    table_before(doc,anchor,'Table S15. Selected paired V–VI contrasts using the common extension bootstrap.',rows,[82,50])
    rows=[['Test fold','V–VI available','Matched/source','Excluded V–VI','Retained (%)']]
    for _,r in coverage.iterrows():rows.append([int(r.heldout_fold),int(r.target_available),int(r.matched_per_source),int(r.target_rows_not_retained),f'{r.target_fraction_retained*100:.1f}'])
    table_before(doc,anchor,'Table S16. Exact diagnosis common support and annotation coverage.',rows,[17,28,30,28,29])
    insert_before(anchor,'Each diagnosis/outcome cell retains min(nV–VI, nlighter) in the non-test pool. The full diagnosis-cell CSV includes cells with zero overlap and all discarded target counts. Every paired panel has identical diagnosis/outcome counts; the held-out test cohort is unchanged.','note')
    rows=[['Training procedure','Mean within-fold AUROC','Fold/panel range']]
    def range_round(x):return str(Decimal(str(x)).quantize(Decimal('0.001'),rounding=ROUND_HALF_EVEN))
    for _,r in fold_summary[fold_summary.cohort=='V-VI'].iterrows():rows.append([LABELS[r.arm],f'{r.mean_within_fold_AUROC:.3f}',f'{range_round(r.minimum_fold_panel_AUROC)} to {range_round(r.maximum_fold_panel_AUROC)}'])
    table_before(doc,anchor,'Table S17. Descriptive post hoc within-fold V–VI ranking diagnostic.',rows,[66,32,34])
    insert_before(anchor,'Mean over the five original folds and, for sampled arms, their 20 panels. The range describes fold/panel estimates; it is not a confidence interval. This analysis was recorded after pooled results were known and uses saved scores without fitting or selection.','note')
    append_node(doc,'Appendix S6. Reproducibility and External-Evaluation Status','h1')
    append_node(doc,'Both original and extension verification start from the pinned feature cache and saved model coefficients, not a separate image re-extraction. The extension uses Python 3.12.14, NumPy 2.3.5, pandas 2.2.3, SciPy 1.17.0, scikit-learn 1.8.0 and threadpoolctl 3.6.0. Scientific plots are regenerated or rendered from their saved numerical/vector sources at 600 dpi; CSV, code and output hashes are retained. The all-data learning-curve position is the observed mean budget 165.6. No image-generation tool or synthetic model results are used.')
    anchor2=doc.add_paragraph('',style=ST['body'])
    items=[('Protocol','plan/MAKE_Exploratory_Extension_20261006_v1.json'),('Held-out scores','predictions/DDI_per_image_predictions.csv.gz; extension_scores.npz'),('Classifier parameters','models/extension_models.npz; extension_models_registry.csv'),('Training row/weight linkage','analysis_image_ids and effective_sample_weights in model arrays'),('Panel membership','splits/adaptation_panels.csv.gz'),('DFR tuning','splits/DFR_inner_folds.csv; results/DFR_inner_tuning.csv'),('DFR validation scores','predictions/DFR_inner_validation_predictions.csv.gz'),('Paired interval samples','uncertainty/paired_bootstrap_samples.npz'),('Independent verification','logs/independent_extension_verification.json; DFR_tuning_reconstruction.json'),('Seven future external heads','models/E4_deployment_models.npz; E4_deployment_registry.csv'),('External evaluator','code/MAKE_Evaluate_MRA_External.py'),('Within-fold diagnostic','Publication package: tables/within_fold_AUROC_draws.csv'),('Figure provenance','Publication package: audit/figure_provenance.json')]
    table_before(doc,anchor2,'Table S18. Traceability inventory (extension archive unless otherwise specified).',[['Evidence','Path or linkage']]+items,[39,93])
    append_node(doc,'The extension contains 1,261 classifier records and 97,744 held-out/transfer score records; a further 91,840 inner-validation predictions were reconstructed from saved heads. The verifier checks fitting exclusions, weights, coefficients, ensemble averages, nested learning panels, exact diagnosis matching and all 216 full-precision paired contrast/interval rows. These counts are computational records and do not increase the number of independently evaluated images.')
    append_node(doc,'E4 has not been performed on real MRA-MIDAS data [6]. The prepared evaluator requires eligible clinical-image features, patient and lesion identifiers, a compatible encoder specification, reviewed diagnosis-to-endpoint mapping and an accepted source-use agreement review. Its proposed lesion-level analysis averages clinical-view logits and uses paired patient-cluster resampling. Synthetic validation fixtures test rejection gates only and are not scientific results. Patient independence across datasets cannot be certified from released DDI identifiers alone.')
    append_node(doc,'References','h1')
    supprefs=[Paragraph(oldrefs[2],doc).text,Paragraph(oldrefs[4],doc).text,Paragraph(oldrefs[5],doc).text,Paragraph(oldrefs[6],doc).text,NEWREFS[16],NEWREFS[32]]
    order=[]
    for q in doc.paragraphs:
        for match in re.findall(r'\[([0-9,]+)\]',q.text):
            for n in map(int,match.split(',')):
                if n not in order:order.append(n)
    assert set(order)==set(range(1,7))
    smap={old:i+1 for i,old in enumerate(order)}
    for q in doc.paragraphs:
        if re.search(r'\[([0-9,]+)\]',q.text):
            set_text(q,re.sub(r'\[([0-9,]+)\]',lambda m:'['+','.join(str(smap[int(n)]) for n in m.group(1).split(','))+']',q.text))
    for n in order:doc.add_paragraph(supprefs[n-1],style=ST['ref'])
    (ROOT/'audit/supplement_reference_inventory.json').write_text(json.dumps([{'number':smap[n],'reference':supprefs[n-1]} for n in order],ensure_ascii=False,indent=2)+'\n')
    update_fields(doc);doc.core_properties.title='Supplementary Materials: '+main.paragraphs[1].text
    doc.save(ROOT/'MAKE_SkinCancer_20261006_Supplementary_MDPI.docx')

if __name__=='__main__':
    make_graphs();main,mapping,oldrefs=build_main();build_supplement(main,mapping,oldrefs)
    (ROOT/'audit/manuscript_source_provenance.json').write_text(json.dumps({'main_source':MAIN.name,'main_source_sha256':sha(MAIN),'supplement_source':SUPP.name,'supplement_source_sha256':sha(SUPP),'source_main_reference_count':15,'final_reference_count':32,'original_tables_1_to_5':'Retained; Table 4 near-zero positive bound shown explicitly','added_main_tables':[6,7,8],'added_main_figures':[5,6,7],'supplement_tables':18,'MRA_MIDAS_performance':'Not performed; no external result fabricated','historical_Table_S13':'Uploaded supplement contains the separately specified verified similarity-group protocol; historical +0.021 result is not included or claimed verified.'},indent=2)+'\n')
    print('Integrated main/supplement written; 32 main references, 8 main tables, 18 supplementary tables.')
