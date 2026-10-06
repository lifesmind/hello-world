"""Check publication tables against full-precision results and rendered PDF text."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
import argparse, hashlib, json, re
import numpy as np
import pandas as pd
from PIL import Image
from docx import Document
import fitz

parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1])
args=parser.parse_args();root=args.root.resolve()
main=Document(root/'MAKE_SkinCancer_20261006_MDPI.docx')
supp=Document(root/'MAKE_SkinCancer_20261006_Supplementary_MDPI.docx')
checks=[]
def check(name,condition,details=None):
    checks.append({'check':name,'pass':bool(condition),'details':details})
def compact(s):return re.sub(r'\s+','',s).replace('\u00ad','').replace('−','-')
def vals(s):return np.array([float(v) for v in re.findall(r'[+−-]?\d+\.\d+',s.replace('−','-'))])
def numcell(name,text,expected,digits=3):
    actual=vals(text);expected=np.round(np.array(expected,float),digits)
    check(name,len(actual)==len(expected) and np.allclose(actual,expected,atol=1e-9,rtol=0),{'display':text,'expected':expected.tolist()})
metrics=pd.read_csv(root/'tables/all_arm_metrics.csv')
contrasts=pd.read_csv(root/'tables/all_paired_contrasts.csv')
v=metrics[metrics.cohort=='V-VI'].set_index('arm')
cv=contrasts[(contrasts.cohort=='V-VI')&(contrasts.metric=='AUROC')]
refit=pd.read_csv(root/'tables/Table_R5_refitting_robustness.csv')
original=pd.read_csv(root/'tables/Table_R3_DDI_arms.csv')
core=pd.read_csv(root/'tables/Table_R4_matched_controls.csv')
transfer=pd.read_csv(root/'tables/Table_R2_PAD_development_DDI_transfer.csv')
check('main scientific table/figure count',len(main.tables)==10 and len(main.inline_shapes)==7)
check('supplement table/figure count',len(supp.tables)==18 and len(supp.inline_shapes)==1)
for i,r in transfer.iterrows():
    numcell(f'Table2 AUROC row{i+1}',main.tables[2].cell(i+1,3).text,[r.AUROC,r.AUROC_CI_low,r.AUROC_CI_high])
for i,arm in enumerate(['PAD_baseline','target_weighted','target_weighted_minus_PAD_baseline'],1):
    r=refit[(refit.stratum=='V-VI')&(refit.metric=='AUROC')&(refit.arm_or_contrast==arm)].iloc[0]
    numcell(f'Table5 original point {arm}',main.tables[5].cell(i,1).text,[r.original_point_estimate])
    numcell(f'Table5 refitted mean {arm}',main.tables[5].cell(i,2).text,[r.bootstrap_mean])
    numcell(f'Table5 refittingCI {arm}',main.tables[5].cell(i,3).text,[r.CI_low,r.CI_high])
arms=['PAD_baseline','target_unweighted','target_weighted','lighter_unweighted','lighter_weighted']
for i,arm in enumerate(arms,1):
    r=original[(original.stratum=='V-VI')&(original.arm==arm)].iloc[0]
    numcell(f'Table3 AUROC {arm}',main.tables[3].cell(i,1).text,[r.AUROC,r.AUROC_CI_low,r.AUROC_CI_high])
    numcell(f'Table3 sensitivity {arm}',main.tables[3].cell(i,2).text,[r.sensitivity_FPR_le_0_50*100],1)
pairs=[('target_weighted','PAD_baseline'),('target_unweighted','PAD_baseline'),('lighter_weighted','PAD_baseline'),('lighter_unweighted','PAD_baseline'),('target_weighted','lighter_weighted'),('target_unweighted','lighter_unweighted'),('target_weighted','target_unweighted')]
for i,(a,b) in enumerate(pairs,1):
    r=core[(core.stratum=='V-VI')&(core.metric=='AUROC')&(core.left_arm==a)&(core.right_arm==b)].iloc[0]
    numcell(f'Table4 change {a} minus {b}',main.tables[4].cell(i,1).text,[r.difference_unrounded])
    if a=='lighter_unweighted':
        actual=vals(main.tables[4].cell(i,2).text)
        check('Table4 positive near-zero lower limit',np.allclose(actual,[round(r.CI_low,5),round(r.CI_high,3)],atol=1e-9,rtol=0))
    else:numcell(f'Table4 CI {a} minus {b}',main.tables[4].cell(i,2).text,[r.CI_low,r.CI_high])
e1=['PAD_baseline','target_weighted','target_unweighted','lighter_weighted','all_DDI_unweighted','all_DDI_source_weighted','all_DDI_group_reweighted','all_DDI_budget_matched','DDI_only_class_balanced','DDI_only_group_balanced','DFR_style']
for i,arm in enumerate(e1,1):
    r=v.loc[arm];numcell(f'Table6 {arm}',main.tables[6].cell(i,2).text,[r.AUROC,r.AUROC_CI_low,r.AUROC_CI_high])
for i,arm in enumerate(['PAD_baseline','target_budget_25','target_budget_50','target_budget_100','target_weighted'],1):
    r=v.loc[arm];numcell(f'Table7 {arm}',main.tables[7].cell(i,1).text,[r.AUROC,r.AUROC_CI_low,r.AUROC_CI_high])
    if arm.startswith('target_budget'):numcell(f'Table7 panel SD {arm}',main.tables[7].cell(i,2).text,[r.AUROC_panel_SD])
for i,arm in enumerate(['diagnosis_matched_target','diagnosis_matched_lighter'],1):
    r=v.loc[arm];numcell(f'Table8 {arm}',main.tables[8].cell(i,1).text,[r.AUROC,r.AUROC_CI_low,r.AUROC_CI_high])
r=cv[(cv.left_arm=='diagnosis_matched_target')&(cv.right_arm=='diagnosis_matched_lighter')].iloc[0]
numcell('Table8 paired diagnosis contrast',main.tables[8].cell(3,1).text,[r.difference_unrounded,r.CI_low,r.CI_high])
rs=pd.read_csv(root/'tables/Table_RS13_proxy_group_changes.csv')
for i,tone in enumerate(['I-II','III-IV','V-VI'],1):
    r=rs[(rs.stratum==tone)&(rs.metric=='AUROC')&(rs.arm=='target_weighted')].iloc[0]
    numcell('TableS13 full precision '+tone,' '.join(supp.tables[12].cell(i,j).text for j in [2,3,4]),[r.baseline,r.adapted,r.difference_unrounded],6)
    numcell('TableS13 CI '+tone,supp.tables[12].cell(i,5).text,[r.CI_low,r.CI_high])
    check('TableS13 subtract before rounding '+tone,abs(r.adapted-r.baseline-r.difference_unrounded)<1e-12)
row=1
for cohort in ['ALL_DDI','I-II','III-IV','V-VI']:
    for _,r in metrics[metrics.cohort==cohort].iterrows():
        numcell(f'TableS14 row {row}',supp.tables[13].cell(row,2).text,[r.AUROC,r.AUROC_CI_low,r.AUROC_CI_high]);row+=1
check('TableS14 has all64 metrics',row==65 and len(supp.tables[13].rows)==65)
pairs15=[('target_weighted','PAD_baseline'),('target_weighted','lighter_weighted'),('target_weighted','all_DDI_unweighted'),('target_weighted','all_DDI_source_weighted'),('target_weighted','all_DDI_group_reweighted'),('target_weighted','all_DDI_budget_matched'),('target_weighted','DDI_only_class_balanced'),('target_weighted','DDI_only_group_balanced'),('target_weighted','DFR_style'),('target_budget_100','target_budget_50'),('target_weighted','target_budget_100'),('diagnosis_matched_target','diagnosis_matched_lighter')]
for i,(a,b) in enumerate(pairs15,1):
    r=cv[(cv.left_arm==a)&(cv.right_arm==b)].iloc[0]
    numcell(f'TableS15 paired contrast row{i}',supp.tables[14].cell(i,1).text,[r.difference_unrounded,r.CI_low,r.CI_high])
folds=pd.read_csv(root/'tables/within_fold_AUROC_summary.csv')
for i,(_,r) in enumerate(folds[folds.cohort=='V-VI'].iterrows(),1):
    numcell(f'TableS17 mean row{i}',supp.tables[16].cell(i,1).text,[r.mean_within_fold_AUROC])
    numcell(f'TableS17 range row{i}',supp.tables[16].cell(i,2).text,[r.minimum_fold_panel_AUROC,r.maximum_fold_panel_AUROC])
coverage=pd.read_csv(root/'tables/Table_E3_matching_coverage.csv')
for i,r in coverage.iterrows():
    cells=supp.tables[15].rows[i+1].cells
    check(f'TableS16 coverage fold{i}',[int(c.text) for c in cells[:4]]==[int(r.heldout_fold),int(r.target_available),int(r.matched_per_source),int(r.target_rows_not_retained)])
check('Figure6 actual all-label budget',coverage.target_available.mean()==165.6 and json.loads((root/'audit/learning_curve_axis_check.json').read_text())['mean_full_budget']==165.6)
# Retained source table texts must match the original inputs.
source_supp=Document(root/'inputs/MAKE_SkinCancer_20261006_Supplementary_MDPI-2_SOURCE.docx')
for i in range(13):
    expected=[[c.text for c in row.cells] for row in source_supp.tables[i].rows]
    actual=[[c.text for c in row.cells] for row in supp.tables[i].rows]
    check(f'Source tableS{i+1} retained',actual==expected)
# All reference identifiers are valid, used, and ordered by first use.
for d,n,name in [(main,32,'main')]:
    refs=[p for p in d.paragraphs if p.style.name=='MDPI_8.1_references']
    order=[]
    for p in d.paragraphs:
        if p.style.name=='MDPI_8.1_references':continue
        for match in re.findall(r'\[([0-9,–-]+)\]',p.text):
            for part in match.split(','):
                ends=re.split('[–-]',part);ids=range(int(ends[0]),int(ends[-1])+1) if len(ends)>1 else [int(part)]
                for v0 in ids:
                    if v0 not in order:order.append(v0)
    check(name+' references complete',len(refs)==n and set(order)==set(range(1,n+1)),{'first_use_order':order,'references':len(refs)})
    check(name+' reference first-use order',order==list(range(1,n+1)))
supp_refs=[p for p in supp.paragraphs if p.style.name=='MDPI_8.1_references']
listed={int(re.match(r'(\d+)\.',p.text).group(1)):re.sub(r'^\d+\.\s*','',p.text) for p in supp_refs}
cited=set()
for p in supp.paragraphs:
    if p.style.name=='MDPI_8.1_references':continue
    for match in re.findall(r'\[([0-9,–-]+)\]',p.text):
        for part in match.split(','):
            ends=re.split('[–-]',part);cited.update(range(int(ends[0]),int(ends[-1])+1))
main_refs=[p.text for p in main.paragraphs if p.style.name=='MDPI_8.1_references']
check('Supplement shares main reference identifiers',set(listed)==cited and len(listed)==6 and all(compact(text)==compact(main_refs[n-1]) for n,text in listed.items()),{'listed':sorted(listed),'cited':sorted(cited)})
check('Abstract length <=200',len(main.paragraphs[6].text.split())<=200)
for name,d in [('MAKE_SkinCancer_20261006_MDPI',main),('MAKE_SkinCancer_20261006_Supplementary_MDPI',supp)]:
    path=root/(name+'.docx')
    with ZipFile(path) as z:
        xml=z.read('word/document.xml')
        check(name+' no tracked changes/comments',all(s not in xml for s in [b'<w:ins ',b'<w:del ',b'commentRangeStart']))
        for i,shape in enumerate(d.inline_shapes,1):
            rid=shape._inline.xpath('.//a:blip/@r:embed')[0];part=d.part.related_parts[rid]
            with Image.open(BytesIO(part.blob)) as im:
                dpi=im.info.get('dpi',(0,0));check(name+f' scientific raster{i} >=600dpi',min(dpi)>=599.9,{'size':im.size,'dpi':dpi})
    pdf=fitz.open(root/(name+'.pdf'));pdftext=compact('\n'.join(p.get_text() for p in pdf))
    failed=[]
    science=d.tables[1:9] if d is main else d.tables
    for ti,t in enumerate(science,1):
        for ri,row in enumerate(t.rows):
            for ci,c in enumerate(row.cells):
                if re.search(r'\d',c.text) and compact(c.text) not in pdftext:failed.append({'table':ti,'row':ri,'col':ci,'text':c.text})
    check(name+' numeric cells present in renderedPDF',not failed,failed)
    bad=[]
    for pi,p in enumerate(pdf,1):
        for word in p.get_text('words'):
            x0,y0,x1,y1=word[:4]
            if x0<0 or y0<0 or x1>p.rect.width+.1 or y1>p.rect.height+.1:bad.append({'page':pi,'word':word[4],'bbox':word[:4]})
    check(name+' text within page boundary',not bad,bad)
    if d is main:check('Abstract on first page','Targeted annotations' in pdf[0].get_text())
    if d is main:check('Publisher copyright fully rendered','CommonsAttribution(CCBY)license' in pdftext)
report={'status':'PASS' if all(c['pass'] for c in checks) else 'FAIL','checks':checks,'check_count':len(checks),'failed_checks':[c for c in checks if not c['pass']]}
(root/'audit/publication_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'status':report['status'],'checks':len(checks),'failures':report['failed_checks']},ensure_ascii=False,indent=2))
raise SystemExit(0 if report['status']=='PASS' else 1)
